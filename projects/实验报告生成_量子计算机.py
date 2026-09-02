#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
金刚石量子计算机实验报告生成脚本
按照示例PDF格式生成Word文档
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def create_quantum_computer_report():
    """创建金刚石量子计算机实验报告"""
    doc = Document()
    
    # 设置页面为Letter尺寸 (8.5 x 11 英寸)
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    
    # 设置默认字体为宋体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 页眉 =====
    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_para.add_run('金刚石量子计算机实验')
    header_run.font.name = '宋体'
    header_run.font.size = Pt(10.5)
    header_run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 中文标题 =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.space_after = Pt(12)
    run = title.add_run('金刚石量子计算机实验')
    run.font.name = '黑体'
    run.font.size = Pt(18)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    # ===== 中文作者 =====
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.space_after = Pt(6)
    run = author.add_run('徐朝阳，张代宣')
    run.font.name = '楷体'
    run.font.size = Pt(14)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
    
    # ===== 中文单位 =====
    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.space_after = Pt(12)
    run = affiliation.add_run('武汉大学物理科学与技术学院，湖北省武汉市 430072')
    run.font.name = '楷体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
    
    # ===== 中文摘要 =====
    abstract_title = doc.add_paragraph()
    abstract_title.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_title.space_before = Pt(6)
    abstract_title.space_after = Pt(3)
    label_run = abstract_title.add_run('摘要：')
    label_run.font.name = '黑体'
    label_run.font.size = Pt(9)
    label_run.font.bold = True
    label_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    abstract_text = doc.add_paragraph()
    abstract_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_text.paragraph_format.first_line_indent = Cm(0.74)
    abstract_text.space_after = Pt(6)
    run = abstract_text.add_run('本实验报告旨在利用金刚石氮-空位色心搭建一个原型量子计算系统，并对其关键性能进行表征与验证。金刚石 NV 色心以其室温下卓越的相干特性、光学初始化和读出能力，成为实现实用化量子计算机极具潜力的物理平台。通过对单个NV 色心进行光谱测量与顺磁共振测量，我们成功实现了电子自旋的初始化、相干操控（拉比振荡）、回波实验 T2、实验动力学去耦和DJ 实验。')
    run.font.name = '宋体'
    run.font.size = Pt(9)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 中文关键词 =====
    keywords = doc.add_paragraph()
    keywords.paragraph_format.first_line_indent = Cm(0.74)
    keywords.space_after = Pt(12)
    label_run = keywords.add_run('关键词：')
    label_run.font.name = '黑体'
    label_run.font.size = Pt(9)
    label_run.font.bold = True
    label_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    kw_run = keywords.add_run('金刚石氮-空位色心；量子计算；量子比特；自旋操控；退相干时间')
    kw_run.font.name = '宋体'
    kw_run.font.size = Pt(9)
    kw_run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 英文标题 =====
    en_title = doc.add_paragraph()
    en_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    en_title.space_before = Pt(12)
    en_title.space_after = Pt(6)
    run = en_title.add_run('Diamond Quantum Computer Experiments')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    
    # ===== 英文作者 =====
    en_author = doc.add_paragraph()
    en_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    en_author.space_after = Pt(6)
    run = en_author.add_run('Xu Chaoyang, Zhang Daixuan')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    
    # ===== 英文单位 =====
    en_affiliation = doc.add_paragraph()
    en_affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    en_affiliation.space_after = Pt(12)
    run = en_affiliation.add_run('School of physical science and technology, Wuhan University, Wuhan, 430072, China')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    
    # ===== 英文摘要 =====
    en_abstract_title = doc.add_paragraph()
    en_abstract_title.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    en_abstract_title.space_before = Pt(6)
    en_abstract_title.space_after = Pt(3)
    label_run = en_abstract_title.add_run('Abstract: ')
    label_run.font.name = 'Times New Roman'
    label_run.font.size = Pt(10.5)
    label_run.font.bold = True
    
    abstract_text = en_abstract_title.add_run('The purpose of this experiment is to construct a prototype quantum computing system using diamond nitrogen-vacancy centers and to characterize and validate its key performance. The diamond NV center, with its exceptional coherence properties at room temperature, along with its capabilities for optical initialization and readout, presents a highly promising physical platform for realizing practical quantum computers. Through spectral measurements and electron spin resonance measurements of a single NV center, we have successfully achieved electron spin initialization, coherent manipulation (Rabi oscillations), spin echo experiments (T2), dynamical decoupling experiments, and the DJ algorithm experiment.')
    abstract_text.font.name = 'Times New Roman'
    abstract_text.font.size = Pt(10.5)
    
    # ===== 英文关键词 =====
    en_keywords = doc.add_paragraph()
    en_keywords.paragraph_format.first_line_indent = Cm(0.74)
    en_keywords.space_after = Pt(12)
    label_run = en_keywords.add_run('Keywords: ')
    label_run.font.name = 'Times New Roman'
    label_run.font.size = Pt(10.5)
    label_run.font.bold = True
    
    kw_run = en_keywords.add_run('Diamond nitrogen-vacancy center (NV center)；Quantum computing；Qubit；Spin manipulation；Decoherence time')
    kw_run.font.name = 'Times New Roman'
    kw_run.font.size = Pt(10.5)
    
    # ===== 引言 =====
    intro_paras = [
        '量子计算是一种遵循量子力学规律调控量子信息单元进行计算的新型计算模式。与经典计算机使用比特（0 或1）作为信息基本单位不同，量子计算机使用量子比特（Qubit）作为基本单位。量子比特可以处于|0>和|1>的叠加态，并通过量子纠缠和量子干涉等特性，在处理特定问题时展现出超越经典计算机的巨大潜力，例如大数分解（Shor 算法）和无序数据库搜索（Grover 算法）。',
        '要实现量子计算，一个物理系统必须满足DiVincenzo 判据，包括：良好的量子比特、初态制备、足够长的相干时间、通用量子逻辑门操作和量子态测量等。金刚石中的氮-空位（Nitrogen-Vacancy, NV）色心因其在室温下具有较长的电子自旋相干时间、可通过光探测磁共振（ODMR）技术进行初始化和读出的特性，成为实现量子计算的一个极具前景的物理平台。',
        '本实验利用基于金刚石NV 色心的量子计算教学机，通过一系列基础实验，旨在理解和掌握量子比特的初始化、操控和读出方法，并最终实现Deutsch-Jozsa 量子算法，直观展示量子计算的并行性优势。'
    ]
    
    for para in intro_paras:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(para)
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 1 实验原理 =====
    section1 = doc.add_paragraph()
    section1.space_before = Pt(12)
    section1.space_after = Pt(6)
    run = section1.add_run('1  实验原理')
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    # 1.1 金刚石NV 色心
    subsection1_1 = doc.add_paragraph()
    subsection1_1.space_before = Pt(6)
    subsection1_1.space_after = Pt(3)
    run = subsection1_1.add_run('1.1  金刚石NV 色心')
    run.font.name = '黑体'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15.6)
    run = p.add_run('NV 色心是金刚石晶格中一个氮原子取代一个碳原子并与之相邻的一个空位所组成的点缺陷。其基态为自旋三重态，在零磁场下，|ms = 0>态与简并的|ms = ±1>态之间存在2.87GHz 的零场劈裂。')
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 1.2 自旋态的初始化与读出
    subsection1_2 = doc.add_paragraph()
    subsection1_2.space_before = Pt(6)
    subsection1_2.space_after = Pt(3)
    run = subsection1_2.add_run('1.2  自旋态的初始化与读出')
    run.font.name = '黑体'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    bullet_items = [
        '初始化：使用532nm 激光照射NV 色心。由于|ms = ±1>态比|ms = 0>态有更高的概率通过系间窜越跃迁到中间单重态，并最终弛豫到|ms = 0>态，经过多个激光激发-弛豫周期后，NV 色心被极化为|ms = 0>态。',
        '读出：|ms = 0>态的荧光强度显著高于|ms = ±1>态。通过探测激光激发下的荧光强度，即可区分量子比特所处的状态。'
    ]
    
    for item in bullet_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run('• ' + item)
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 1.3 自旋态的操控
    subsection1_3 = doc.add_paragraph()
    subsection1_3.space_before = Pt(6)
    subsection1_3.space_after = Pt(3)
    run = subsection1_3.add_run('1.3  自旋态的操控')
    run.font.name = '黑体'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15.6)
    run = p.add_run('通过施加与自旋能级差共振的微波场，可以实现对电子自旋态的相干操控。在共振条件下（微波频率ω = ω0），量子态在|0>和|1>之间以拉比频率ω1 做周期性振荡，其演化规律由薛定谔方程描述。')
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    bullet_items2 = [
        'π 脉冲：当微波作用时间满足ω1t = π 时，可实现|0><->|1>的完全翻转，对应量子非门操作。',
        'π/2 脉冲：当微波作用时间满足ω1t = π/2 时，可将量子比特从本征态制备到叠加态。'
    ]
    
    for item in bullet_items2:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run('• ' + item)
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 1.4 Deutsch-Jozsa 算法
    subsection1_4 = doc.add_paragraph()
    subsection1_4.space_before = Pt(6)
    subsection1_4.space_after = Pt(3)
    run = subsection1_4.add_run('1.4  Deutsch-Jozsa 算法')
    run.font.name = '黑体'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15.6)
    run = p.add_run('D-J 算法用于判断一个函数是常函数（所有输入输出相同）还是平衡函数（一半输入输出0，另一半输出1）。对于n = 1 的情况，经典算法最坏需要2 次函数查询，而D-J 量子算法仅需1 次。其核心在于利用量子叠加性和相因子，通过量子线路的操作，最终通过对第一个量子比特的测量即可确定性地地区分函数类型。')
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 2 实验装置 =====
    section2 = doc.add_paragraph()
    section2.space_before = Pt(12)
    section2.space_after = Pt(6)
    run = section2.add_run('2  实验装置')
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15.6)
    run = p.add_run('实验使用"金刚石量子计算教学机"，其主要由以下模块构成：')
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    device_items = [
        '光学模块：产生532nm 激光用于初始化和读出；通过共聚焦光路收集NV 色心的荧光并由光电探测器转换为电信号。',
        '微波模块：产生频率和功率可调的微波信号，经放大和脉冲调制后，通过天线辐射至NV 色心，用于操控自旋态。',
        '控制与采集模块：产生精确定时的TTL 控制脉冲，同步激光器、微波开关和数据采集卡的时序。',
        '软件系统：通过Diamond I Studio 软件控制整个实验流程。'
    ]
    
    for item in device_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run('• ' + item)
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 3 实验内容与步骤 =====
    section3 = doc.add_paragraph()
    section3.space_before = Pt(12)
    section3.space_after = Pt(6)
    run = section3.add_run('3  实验内容与步骤')
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    experiments = [
        {
            'title': '3.1  连续波ODMR 实验',
            'purpose': '目的：测量NV 色心的光探测磁共振谱，确定其自旋共振频率。',
            'steps': '步骤：设置微波频率扫描范围，在连续激光照射和连续微波扫描下，采集荧光信号。',
            'fig': '图1: 连续波ODMR 谱',
            'result': '结果：测得两个共振频率MW1 = 2845 MHz，MW2 = 2898 MHz。'
        },
        {
            'title': '3.2  拉比振荡实验',
            'purpose': '目的：验证对量子比特的相干操控，并标定实现量子逻辑门所需的微波脉冲宽度。',
            'steps': '步骤：固定微波频率于一个共振频率，扫描微波脉冲的宽度，测量荧光信号。',
            'fig': '图2: 拉比振荡曲线，MW1\n图3: 拉比振荡曲线，MW2',
            'result': '结果：从拉比振荡曲线中，在MW1 波源下，测得π/2 脉冲宽度为110 ns，π 脉冲宽度为230 ns。在MW2 波源下，测得π/2 脉冲宽度为150 ns，π 脉冲宽度为330 ns。'
        },
        {
            'title': '3.3  回波实验',
            'purpose': '目的：演示使用Hahn 回波序列来抑制低频噪声的影响。',
            'steps': '步骤：执行(π/2)−τ−π−τ−(π/2)的脉冲序列，扫描第二个时间间隔τ。',
            'fig': '图4: 回波信号，MW1\n图5: 回波信号，MW2',
            'result': '结果：观察到了回波现象，表明成功应用了动态解耦技术。'
        },
        {
            'title': '3.4  T2 退相干时间测量实验',
            'purpose': '目的：定量测量NV 色心电子自旋的退相干时间T2。',
            'steps': '步骤：在Hahn 回波序列中，扫描总自由演化时间τ，测量回波幅度的衰减曲线。',
            'fig': '图6: T2 衰减曲线及指数拟合',
            'result': '结果：通过拟合衰减曲线，得到退相干时间T2 = _______ μs。'
        },
        {
            'title': '3.5  D-J 算法实验',
            'purpose': '目的：在NV 色心系统上实验实现一阶Deutsch-Jozsa 算法。',
            'steps': '步骤：通过特定的微波脉冲序列来构造四个不同的Uf Oracle 门。',
            'fig': '图7: D-J1 算法实验结果\n图8: D-J2 算法实验结果\n图9: D-J3 算法实验结果\n图10: D-J4 算法实验结果',
            'result': '结果：\n• 对于序列DJ1 和DJ2，回波信号方向如图，判断为常函数。\n• 对于序列DJ3 和DJ4，回波信号方向为如图，判断为平衡函数。\n实验结果与理论预言一致，成功演示了D-J 算法。'
        }
    ]
    
    for exp in experiments:
        # 小节标题
        p = doc.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(3)
        run = p.add_run(exp['title'])
        run.font.name = '黑体'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        
        # 目的
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(exp['purpose'])
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 步骤
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(exp['steps'])
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 图注
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(exp['fig'])
        run.font.name = '宋体'
        run.font.size = Pt(9)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 结果
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(exp['result'])
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 4 讨论与分析 =====
    section4 = doc.add_paragraph()
    section4.space_before = Pt(12)
    section4.space_after = Pt(6)
    run = section4.add_run('4  讨论与分析')
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    discussion_items = [
        '1. 量子态操控：拉比振荡的成功观测，直接证明了我们能够对NV 色心自旋进行相干的量子态操控。',
        '2. 退相干：T2 实验结果表明量子叠加态会随着时间演化而衰减，体现了环境噪声对量子系统的影响。',
        '3. 量子优势：D-J 算法实验以最直观的方式证明，对于特定问题，量子算法可以通过一次查询解决经典算法需要两次查询才能确定的问题。'
    ]
    
    for item in discussion_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(item)
        run.font.name = '宋体'
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 5 结论 =====
    section5 = doc.add_paragraph()
    section5.space_before = Pt(12)
    section5.space_after = Pt(6)
    run = section5.add_run('5  结论')
    run.font.name = '黑体'
    run.font.size = Pt(12)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15.6)
    run = p.add_run('本实验成功利用金刚石NV 色心系统，完成了从量子比特性质表征、基本操作标定、相干特性研究到简单量子算法演示的全过程。实验结果与理论预期相符，加深了我们对量子计算基本原理和实验实现方法的理解。金刚石NV 色心作为一种优秀的室温固态量子比特平台，在量子计算教学和研究中具有重要价值。')
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # ===== 参考文献 =====
    references = doc.add_paragraph()
    references.space_before = Pt(12)
    references.space_after = Pt(6)
    run = references.add_run('参考文献')
    run.font.name = '黑体'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    
    refs = [
        '[1] D\'Aloia A G, et al. Characterization of a commercial diamond sensor under alpha-particle irradiation. Diamond and Related Materials, 2023.',
        '[2] Kizil O, et al. Nitrogen-vacancy centers in diamond: A review. Diamond and Related Materials, 2022.',
        '[3] DiVincenzo D P. The physical implementation of quantumcomputing. Fortschritte der Physik, 2000, 48(9-11): 771-783.'
    ]
    
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = Pt(15.6)
        run = p.add_run(ref)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)
    
    # 保存文档
    output_path = '实例实验报告/金刚石量子计算机实验报告.docx'
    doc.save(output_path)
    print(f'✓ 已生成：{output_path}')
    return output_path


if __name__ == '__main__':
    output_path = create_quantum_computer_report()
    print(f'报告已保存至：{output_path}')
