from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_TEACHING_MATERIAL_BYTES = 20 * 1024 * 1024
MAX_TEACHING_MATERIAL_TEXT = 120_000
TEACHING_MATERIAL_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


def extract_teaching_material(upload) -> tuple[str, str]:
    """Extract bounded text from a teacher-provided experiment reference file."""
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix == ".doc":
        raise ValueError("旧版 DOC 暂不支持，请另存为 DOCX 后上传")
    if suffix not in TEACHING_MATERIAL_SUFFIXES:
        raise ValueError("实验材料仅支持 PDF、DOCX、TXT 或 Markdown")
    data = upload.stream.read(MAX_TEACHING_MATERIAL_BYTES + 1)
    upload.stream.seek(0)
    if not data:
        raise ValueError("上传的实验材料为空")
    if len(data) > MAX_TEACHING_MATERIAL_BYTES:
        raise ValueError("实验材料不能超过 20 MB")
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("加密 PDF 无法读取，请上传未加密版本")
            parts = [(page.extract_text() or "").strip() for page in reader.pages[:300]]
            text = "\n\n".join(part for part in parts if part)
        elif suffix == ".docx":
            document = Document(io.BytesIO(data))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if line:
                        parts.append(line)
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("实验材料解析失败，请确认文件没有损坏") from exc
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 20:
        raise ValueError("材料中没有提取到足够文字；扫描版 PDF 请先进行 OCR")
    return text[:MAX_TEACHING_MATERIAL_TEXT], suffix.removeprefix(".")
