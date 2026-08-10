from __future__ import annotations

import base64
import io
import json
import re
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from flask import current_app
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..db import db_session
from ..models import FileAsset
from .report_drafts import document_from_payload


_FONT_NAME = "WHUSimHei"
_INK = "#111111"
_MUTED = "#4B5563"
_RULE = "#202020"
_SUPERSCRIPT_ASCII = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾", "0123456789+-=()")
_SUBSCRIPT_ASCII = str.maketrans("₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎", "0123456789+-=()")
_LATEX_SCRIPT = str.maketrans({
    **dict(zip("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")),
})
_LATEX_SUPER = str.maketrans({
    **dict(zip("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")),
})


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    if isinstance(value, str):
        return value.strip() or fallback
    return str(value)


def _latex_plain(value: Any) -> str:
    """Render the supported LaTeX subset as portable Unicode for DOCX."""
    text = _text(value)
    replacements = {r"\Omega": "Ω", r"\omega": "ω", r"\pm": "±", r"\times": "×", r"\cdot": "·", r"\le": "≤", r"\ge": "≥"}
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = text.replace(r"\left", "").replace(r"\right", "")
    text = re.sub(r"_\{([^{}]+)\}", lambda match: match.group(1).translate(_LATEX_SCRIPT), text)
    text = re.sub(r"\^\{([^{}]+)\}", lambda match: match.group(1).translate(_LATEX_SUPER), text)
    text = re.sub(r"_([0-9+\-=()])", lambda match: match.group(1).translate(_LATEX_SCRIPT), text)
    text = re.sub(r"\^([0-9+\-=()])", lambda match: match.group(1).translate(_LATEX_SUPER), text)
    text = re.sub(r"\\([A-Za-z]+)", lambda match: match.group(1), text)
    return text.replace("{", "").replace("}", "")


def _latex_pdf_markup(value: Any) -> str:
    """Use ReportLab super/sub markup so the PDF never depends on rare glyphs."""
    text = _latex_plain(value)
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    for character, ascii_value in zip("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾", "0123456789+-=()"):
        escaped = escaped.replace(character, f"<super>{ascii_value}</super>")
    for character, ascii_value in zip("₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎", "0123456789+-=()"):
        escaped = escaped.replace(character, f"<sub>{ascii_value}</sub>")
    return escaped


def _asset_bytes(asset_hash: str) -> bytes | None:
    asset = db_session().get(FileAsset, asset_hash)
    if not asset or not asset.content_type.startswith("image/"):
        return None
    path = Path(current_app.config["DATA_DIR"]) / asset.object_key
    try:
        return path.read_bytes() if path.is_file() else None
    except OSError:
        return None


def _plot_bytes(document: Mapping[str, Any]) -> bytes | None:
    value = _mapping(document.get("fit_result")).get("plot_data_url", "")
    if not isinstance(value, str) or not value.startswith("data:image/") or "," not in value:
        return None
    try:
        return base64.b64decode(value.split(",", 1)[1], validate=True)
    except Exception:
        return None


def report_document(revision, version, exp, user) -> dict:
    return document_from_payload(_mapping(getattr(revision, "payload", {})), version, exp, user)


def _report_context(revision, version, exp, evaluation, user, review=None, teacher_view=False) -> dict:
    document = report_document(revision, version, exp, user)
    definition = _mapping(getattr(version, "definition", {}))
    return {
        "document": document,
        "title": document.get("title_zh") or getattr(exp, "title", "实验"),
        "student_name": getattr(user, "name", ""),
        "student_id": getattr(user, "id", ""),
        "version_no": getattr(version, "version_no", "-"),
        "revision_no": getattr(revision, "revision_no", "-"),
        "submitted_at": getattr(revision, "submitted_at", None),
        "reference_value": _text(definition.get("reference_value")),
        "reference_unit": _text(definition.get("reference_unit")),
        "evaluation": evaluation,
        "teacher_view": bool(teacher_view),
        "review": review,
    }


def _set_run_font(run, name: str, size: float, *, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), "宋体" if name == "Times New Roman" else name)


def _set_cell_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, values in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in values.items():
            element.set(qn(f"w:{key}"), str(value))


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    _set_run_font(run, "Times New Roman", 9)


def _configure_docx(document: Document, title: str):
    section = document.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Mm(20); section.bottom_margin = Mm(20)
    section.left_margin = Mm(23); section.right_margin = Mm(23)
    section.header_distance = Mm(8); section.footer_distance = Mm(8)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"; normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(2)
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Mm(164))
    table.autofit = False
    table.columns[0].width = Mm(140); table.columns[1].width = Mm(24)
    left, right = table.rows[0].cells
    left.width = Mm(140); right.width = Mm(24)
    for cell in (left, right):
        _set_cell_borders(cell, bottom={"val": "single", "sz": "8", "color": "202020"})
    p = left.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(p.add_run(title), "宋体", 9)
    p = right.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _add_page_field(p)


def _add_docx_paragraph(document: Document, text: str, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=10.5, indent=True):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(2)
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    run = paragraph.add_run(text or "")
    _set_run_font(run, "宋体", size, bold=bold, italic=italic)
    return paragraph


def _add_docx_rich_paragraph(document: Document, block: Mapping[str, Any], *, align, kind: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = align; paragraph.paragraph_format.line_spacing = 1.25; paragraph.paragraph_format.space_after = Pt(2)
    if kind == "paragraph": paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if kind == "quote": paragraph.paragraph_format.left_indent = Mm(8)
    runs = block.get("runs") if isinstance(block.get("runs"), list) else None
    for item in runs or [{"text": block.get("text", ""), "marks": []}]:
        marks = set(item.get("marks") or [])
        run = paragraph.add_run(_text(item.get("text")))
        _set_run_font(run, "宋体", 10.5, bold=kind == "heading" or "bold" in marks, italic=kind == "quote" or "italic" in marks)
        run.font.underline = "underline" in marks
        run.font.superscript = "superscript" in marks
        run.font.subscript = "subscript" in marks
    return paragraph


def _add_docx_labeled(document: Document, label: str, text: str, *, english=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.18
    paragraph.paragraph_format.space_after = Pt(2)
    _set_run_font(paragraph.add_run(label), "Times New Roman" if english else "宋体", 9.5, bold=True)
    _set_run_font(paragraph.add_run(text or ""), "Times New Roman" if english else "宋体", 9.5)


def _add_docx_table(document: Document, rows: list[list[str]], caption: str = ""):
    if not rows:
        return
    if caption:
        p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(caption), "宋体", 9)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = True
    for row_index, values in enumerate(rows):
        if row_index == 0:
            _set_repeat_table_header(table.rows[row_index])
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            _set_run_font(p.add_run(str(value)), "宋体", 9, bold=row_index == 0)
            edges = {"left": {"val": "nil"}, "right": {"val": "nil"}}
            if row_index == 0:
                edges.update(top={"val": "single", "sz": "10", "color": "202020"}, bottom={"val": "single", "sz": "6", "color": "202020"})
            if row_index == len(rows) - 1:
                edges["bottom"] = {"val": "single", "sz": "10", "color": "202020"}
            _set_cell_borders(cell, **edges)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_docx_image(document: Document, data: bytes, width_percent: int, caption: str):
    paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        paragraph.add_run().add_picture(io.BytesIO(data), width=Mm(155 * width_percent / 100))
    except Exception:
        _set_run_font(paragraph.add_run("[图片无法读取]"), "宋体", 9)
    if caption:
        p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(caption), "宋体", 9)


def _add_docx_blocks(document: Document, blocks: list[dict]):
    for block in blocks:
        kind = block.get("type")
        if kind in {"paragraph", "heading", "quote"}:
            align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(block.get("alignment"), WD_ALIGN_PARAGRAPH.JUSTIFY)
            _add_docx_rich_paragraph(document, block, align=align, kind=kind)
        elif kind == "formula":
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(p.add_run(_latex_plain(block.get("latex"))), "Times New Roman", 11, italic=True)
            if block.get("caption"):
                _add_docx_paragraph(document, _text(block.get("caption")), align=WD_ALIGN_PARAGRAPH.CENTER, size=9, indent=False)
        elif kind == "list":
            for index, item in enumerate(block.get("items", []), 1):
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Mm(6); p.paragraph_format.first_line_indent = Mm(-6); p.paragraph_format.space_after = Pt(1)
                prefix = f"{index}.  " if block.get("ordered") else "•  "
                _set_run_font(p.add_run(prefix + _text(item)), "宋体", 10.5)
        elif kind == "table":
            _add_docx_table(document, [[_text(cell) for cell in row] for row in block.get("rows", [])], _text(block.get("caption")))
        elif kind == "image":
            data = _asset_bytes(_text(block.get("asset_hash")))
            if data:
                _add_docx_image(document, data, int(block.get("width", 80)), _text(block.get("caption")))


def build_docx(revision, version, exp, evaluation, user, review=None, teacher_view=False) -> io.BytesIO:
    data = _report_context(revision, version, exp, evaluation, user, review, teacher_view)
    report = data["document"]
    document = Document()
    _configure_docx(document, data["title"])

    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    _set_run_font(p.add_run(report["title_zh"]), "宋体", 18, bold=True)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
    _set_run_font(p.add_run(f"{report['author_name']}  ({report['student_id']})"), "宋体", 10.5)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    _set_run_font(p.add_run(report["affiliation_zh"]), "宋体", 9.5)
    _add_docx_labeled(document, "摘要：", report["abstract_zh"])
    _add_docx_labeled(document, "关键词：", report["keywords_zh"])

    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    _set_run_font(p.add_run(report["title_en"]), "Times New Roman", 15, bold=True)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
    _set_run_font(p.add_run(f"{report['author_name']}  ({report['student_id']})"), "Times New Roman", 10)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(5)
    _set_run_font(p.add_run(report["affiliation_en"]), "Times New Roman", 9.5)
    _add_docx_labeled(document, "Abstract: ", report["abstract_en"], english=True)
    _add_docx_labeled(document, "Keywords: ", report["keywords_en"], english=True)

    for index, section in enumerate(report.get("sections", []), 1):
        p = document.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)
        _set_run_font(p.add_run(f"{index}  {section['title']}"), "宋体", 12, bold=True)
        _add_docx_blocks(document, section.get("blocks", []))
        if section.get("id") == "fit":
            plot = _plot_bytes(report)
            if plot:
                _add_docx_image(document, plot, 88, "实验数据与拟合结果")

    if teacher_view:
        document.add_section(WD_SECTION.NEW_PAGE)
        p = document.add_paragraph(); _set_run_font(p.add_run("教师内部审核"), "宋体", 14, bold=True)
        _add_docx_paragraph(document, "内部资料，仅限教师查看，不属于学生报告正文。", indent=False)
        _add_docx_table(document, [
            ["项目", "内容"],
            ["人工评分", _text(getattr(review, "private_score", None), "尚未审核")],
            ["内部评语", _text(getattr(review, "private_comment", None), "尚未填写")],
            ["报告修订", f"R{data['revision_no']}"],
        ])

    stream = io.BytesIO(); document.save(stream); stream.seek(0); return stream


def _pdf_text(value: Any) -> str:
    text = _text(value)
    text = text.translate(_SUPERSCRIPT_ASCII).translate(_SUBSCRIPT_ASCII)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def _pdf_inline(block: Mapping[str, Any]) -> str:
    runs = block.get("runs") if isinstance(block.get("runs"), list) else None
    if not runs:
        return _pdf_text(block.get("text"))
    rendered = []
    for item in runs:
        value = _pdf_text(item.get("text")); marks = set(item.get("marks") or [])
        if "bold" in marks: value = f"<b>{value}</b>"
        if "italic" in marks: value = f"<i>{value}</i>"
        if "underline" in marks: value = f"<u>{value}</u>"
        if "superscript" in marks: value = f"<super>{value}</super>"
        if "subscript" in marks: value = f"<sub>{value}</sub>"
        rendered.append(value)
    return "".join(rendered)


def _pdf_styles():
    base = getSampleStyleSheet()
    return {
        "zh_title": ParagraphStyle("zh_title", parent=base["Title"], fontName=_FONT_NAME, fontSize=18, leading=23, alignment=TA_CENTER, textColor=colors.HexColor(_INK), spaceAfter=3 * mm),
        "en_title": ParagraphStyle("en_title", parent=base["Title"], fontName="Times-Bold", fontSize=15, leading=19, alignment=TA_CENTER, textColor=colors.HexColor(_INK), spaceBefore=4 * mm, spaceAfter=2 * mm),
        "center": ParagraphStyle("center", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=9.5, leading=13, alignment=TA_CENTER, spaceAfter=1.5 * mm),
        "en_center": ParagraphStyle("en_center", parent=base["BodyText"], fontName="Times-Roman", fontSize=9.5, leading=12, alignment=TA_CENTER, spaceAfter=1.2 * mm),
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=9.5, leading=14.5, alignment=TA_JUSTIFY, firstLineIndent=7 * mm, spaceAfter=1.4 * mm, wordWrap="CJK"),
        "en_body": ParagraphStyle("en_body", parent=base["BodyText"], fontName="Times-Roman", fontSize=9, leading=13, alignment=TA_JUSTIFY, spaceAfter=1.4 * mm),
        "heading": ParagraphStyle("heading", parent=base["Heading1"], fontName=_FONT_NAME, fontSize=11.5, leading=15, textColor=colors.HexColor(_INK), spaceBefore=4 * mm, spaceAfter=1.5 * mm, keepWithNext=True),
        "formula": ParagraphStyle("formula", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=11, leading=15, alignment=TA_CENTER, spaceBefore=1.5 * mm, spaceAfter=1.5 * mm),
        "caption": ParagraphStyle("caption", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=8.5, leading=11, alignment=TA_CENTER, spaceAfter=2 * mm),
        "quote": ParagraphStyle("quote", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=9.3, leading=14, leftIndent=7 * mm, rightIndent=7 * mm, textColor=colors.HexColor(_MUTED), spaceAfter=2 * mm),
        "small": ParagraphStyle("small", parent=base["BodyText"], fontName=_FONT_NAME, fontSize=8.5, leading=11, alignment=TA_LEFT),
    }


def _pdf_labeled(label: str, text: str, styles, english=False):
    style = styles["en_body" if english else "body"]
    style = ParagraphStyle(f"label-{label}-{english}", parent=style, firstLineIndent=0)
    return Paragraph(f"<b>{_pdf_text(label)}</b>{_pdf_text(text)}", style)


def _pdf_table(rows: list[list[str]], styles, caption: str = ""):
    flowables = []
    if caption:
        flowables.append(Paragraph(_pdf_text(caption), styles["caption"]))
    if not rows:
        return flowables
    count = len(rows[0]); available = 158 * mm
    widths = [available / count] * count
    cells = [[Paragraph(_pdf_text(cell), styles["small"]) for cell in row] for row in rows]
    table = Table(cells, colWidths=widths, repeatRows=1, hAlign="CENTER")
    commands = [
        ("FONTNAME", (0, 0), (-1, -1), _FONT_NAME), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LINEABOVE", (0, 0), (-1, 0), 0.9, colors.HexColor(_RULE)),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor(_RULE)),
        ("LINEBELOW", (0, -1), (-1, -1), 0.9, colors.HexColor(_RULE)),
    ]
    table.setStyle(TableStyle(commands)); flowables.append(table); flowables.append(Spacer(1, 2 * mm))
    return flowables


def _pdf_image(data: bytes, width_percent: int):
    image = Image(io.BytesIO(data))
    max_width, max_height = 150 * mm * width_percent / 100, 95 * mm
    scale = min(max_width / image.imageWidth, max_height / image.imageHeight, 1)
    image.drawWidth = image.imageWidth * scale; image.drawHeight = image.imageHeight * scale; image.hAlign = "CENTER"
    return image


def _pdf_blocks(blocks: list[dict], styles) -> list:
    story: list = []
    for block in blocks:
        kind = block.get("type")
        if kind in {"paragraph", "heading", "quote"}:
            style = styles["quote"] if kind == "quote" else styles["body"]
            if kind == "heading":
                style = ParagraphStyle("block_heading", parent=styles["body"], firstLineIndent=0, fontSize=10, leading=14, spaceBefore=1 * mm, keepWithNext=True)
            story.append(Paragraph(_pdf_inline(block), style))
        elif kind == "formula":
            formula = Paragraph(_latex_pdf_markup(block.get("latex")), styles["formula"])
            caption = Paragraph(_pdf_text(block.get("caption")), styles["caption"]) if block.get("caption") else None
            story.append(KeepTogether([formula, caption] if caption else [formula]))
        elif kind == "list":
            items = [ListItem(Paragraph(_pdf_text(item), ParagraphStyle("list_body", parent=styles["body"], firstLineIndent=0)), leftIndent=4 * mm) for item in block.get("items", [])]
            if items:
                story.append(ListFlowable(items, bulletType="1" if block.get("ordered") else "bullet", leftIndent=8 * mm, bulletFontName=_FONT_NAME, bulletFontSize=8.5))
        elif kind == "table":
            story.extend(_pdf_table([[_text(cell) for cell in row] for row in block.get("rows", [])], styles, _text(block.get("caption"))))
        elif kind == "image":
            data = _asset_bytes(_text(block.get("asset_hash")))
            if data:
                image = _pdf_image(data, int(block.get("width", 80)))
                caption = Paragraph(_pdf_text(block.get("caption")), styles["caption"]) if block.get("caption") else None
                story.append(KeepTogether([image, caption] if caption else [image]))
    return story


def build_pdf(revision, version, exp, evaluation, user, font_path: Path, review=None, teacher_view=False) -> io.BytesIO:
    pdfmetrics.registerFont(TTFont(_FONT_NAME, str(font_path)))
    data = _report_context(revision, version, exp, evaluation, user, review, teacher_view)
    report = data["document"]; styles = _pdf_styles(); buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=23 * mm, rightMargin=23 * mm, topMargin=21 * mm, bottomMargin=19 * mm, title=f"{data['title']}实验报告", author="武汉大学物理科学与技术学院", subject="正式实验报告", allowSplitting=True)
    story: list = [
        Paragraph(_pdf_text(report["title_zh"]), styles["zh_title"]),
        Paragraph(_pdf_text(f"{report['author_name']}  ({report['student_id']})"), styles["center"]),
        Paragraph(_pdf_text(report["affiliation_zh"]), styles["center"]),
        _pdf_labeled("摘要：", report["abstract_zh"], styles),
        _pdf_labeled("关键词：", report["keywords_zh"], styles),
        Paragraph(_pdf_text(report["title_en"]), styles["en_title"]),
        Paragraph(_pdf_text(f"{report['author_name']}  ({report['student_id']})"), styles["en_center"]),
        Paragraph(_pdf_text(report["affiliation_en"]), styles["en_center"]),
        _pdf_labeled("Abstract: ", report["abstract_en"], styles, english=True),
        _pdf_labeled("Keywords: ", report["keywords_en"], styles, english=True),
    ]
    for index, section in enumerate(report.get("sections", []), 1):
        story.append(Paragraph(_pdf_text(f"{index}  {section['title']}"), styles["heading"]))
        story.extend(_pdf_blocks(section.get("blocks", []), styles))
        if section.get("id") == "fit":
            plot = _plot_bytes(report)
            if plot:
                story.append(KeepTogether([_pdf_image(plot, 88), Paragraph("实验数据与拟合结果", styles["caption"])]))
    if teacher_view:
        story.append(Paragraph("教师内部审核", styles["heading"]))
        story.append(Paragraph("内部资料，仅限教师查看，不属于学生报告正文。", styles["body"]))
        story.extend(_pdf_table([
            ["项目", "内容"],
            ["人工评分", _text(getattr(review, "private_score", None), "尚未审核")],
            ["内部评语", _text(getattr(review, "private_comment", None), "尚未填写")],
            ["报告修订", f"R{data['revision_no']}"],
        ], styles))

    def header_footer(canvas, _doc):
        canvas.saveState(); canvas.setStrokeColor(colors.HexColor(_RULE)); canvas.setLineWidth(0.5)
        canvas.line(23 * mm, A4[1] - 13 * mm, A4[0] - 23 * mm, A4[1] - 13 * mm)
        canvas.setFont(_FONT_NAME, 8.5); canvas.setFillColor(colors.HexColor(_INK))
        canvas.drawString(23 * mm, A4[1] - 10.5 * mm, data["title"])
        canvas.drawRightString(A4[0] - 23 * mm, A4[1] - 10.5 * mm, str(canvas.getPageNumber()))
        canvas.restoreState()

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer); buffer.seek(0); return buffer
