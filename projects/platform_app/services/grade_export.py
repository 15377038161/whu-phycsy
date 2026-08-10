from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from ..models import Course, CourseExperiment, Enrollment, Experiment, ExperimentVersion, Review, SubmissionRevision, User


def build_grade_workbook(session, course: Course) -> io.BytesIO:
    students = session.execute(
        select(User)
        .join(Enrollment, Enrollment.student_id == User.id)
        .where(Enrollment.course_id == course.id, User.role == "student")
        .order_by(User.id)
    ).scalars().all()
    experiments = session.execute(
        select(Experiment)
        .join(CourseExperiment, CourseExperiment.experiment_id == Experiment.id)
        .where(CourseExperiment.course_id == course.id, CourseExperiment.active.is_(True))
        .order_by(CourseExperiment.position, Experiment.order_no, Experiment.id)
    ).scalars().all()
    student_ids = [student.id for student in students]
    experiment_ids = [experiment.id for experiment in experiments]
    revision_rows = session.execute(
        select(SubmissionRevision, ExperimentVersion)
        .join(ExperimentVersion, ExperimentVersion.id == SubmissionRevision.experiment_version_id)
        .where(
            SubmissionRevision.course_id == course.id,
            SubmissionRevision.student_id.in_(student_ids or ["__none__"]),
            ExperimentVersion.experiment_id.in_(experiment_ids or ["__none__"]),
            SubmissionRevision.status == "submitted",
        )
        .order_by(SubmissionRevision.submitted_at.desc(), SubmissionRevision.revision_no.desc())
    ).all()
    latest: dict[tuple[str, str], SubmissionRevision] = {}
    for revision, version in revision_rows:
        latest.setdefault((revision.student_id, version.experiment_id), revision)
    revision_ids = [revision.id for revision in latest.values()]
    review_rows = session.scalars(
        select(Review)
        .where(Review.submission_id.in_(revision_ids or ["__none__"]), Review.superseded.is_(False))
        .order_by(Review.created_at.desc())
    ).all()
    reviews: dict[str, Review] = {}
    for review in review_rows:
        reviews.setdefault(review.submission_id, review)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "实验成绩"
    end_column = 2 + len(experiments) * 2 + 4
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet["A1"] = f"{course.name} 实验成绩表"
    sheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="123B56")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    sheet["A2"] = f"导出时间：{datetime.now(timezone.utc).astimezone().strftime('%Y-%m-%d %H:%M')}　成绩口径：教师最新有效人工审核；空白表示尚无有效成绩"
    sheet["A2"].font = Font(name="微软雅黑", size=9, color="666666")

    headers = ["学号", "姓名"]
    for experiment in experiments:
        headers.extend([f"{experiment.title}成绩", f"{experiment.title}审核状态"])
    headers.extend(["已评分实验数", "总分", "平均分", "备注"])
    header_row = 4
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287F7A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="D9E1E5")
    for row_index, student in enumerate(students, header_row + 1):
        values: list[object] = [student.id, student.name]
        scores: list[int] = []
        for experiment in experiments:
            revision = latest.get((student.id, experiment.id))
            review = reviews.get(revision.id) if revision else None
            if not revision:
                values.extend([None, "未提交"])
            elif not review:
                values.extend([None, f"R{revision.revision_no} 待审核"])
            elif review.decision == "approved":
                values.extend([review.private_score, f"R{revision.revision_no} 已评分"])
                scores.append(review.private_score)
            elif review.decision == "returned":
                values.extend([None, f"R{revision.revision_no} 已退回"])
            else:
                values.extend([None, f"R{revision.revision_no} 审核草稿"])
        values.extend([len(scores), sum(scores) if scores else None, round(sum(scores) / len(scores), 2) if scores else None, ""])
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F7F6")

    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(end_column)}{max(header_row, header_row + len(students))}"
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 14
    for column in range(3, end_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 15 if column <= 2 + len(experiments) * 2 else 13
    sheet.sheet_view.showGridLines = False
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
