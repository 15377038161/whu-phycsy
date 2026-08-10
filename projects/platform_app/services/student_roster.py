from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path


MAX_ROSTER_BYTES = 10 * 1024 * 1024
MAX_ROSTER_ROWS = 5_000
MAX_ROSTER_COLUMNS = 80
HEADER_SCAN_ROWS = 30


def _normalized_header(value: str) -> str:
    return re.sub(r"[\s_＿\-—()（）\[\]【】/:：.]+", "", str(value or "").strip().lower())


ID_ALIASES = {_normalized_header(value) for value in (
    "学号", "学生学号", "学籍号", "学生编号", "账号", "登录账号", "用户名",
    "一卡通号", "校园卡号", "student id", "student number", "student no",
    "studentid", "studentno", "userid", "loginid", "sno", "sid",
)}
NAME_ALIASES = {_normalized_header(value) for value in (
    "姓名", "学生姓名", "名字", "中文姓名", "姓名中文", "student name", "studentname", "name",
)}
PASSWORD_ALIASES = {_normalized_header(value) for value in (
    "临时密码", "初始密码", "初始登录密码", "登录密码", "默认密码", "密码",
    "temporary password", "initial password", "password", "passwd", "pwd",
)}
ID_NEGATIVE_HEADERS = {_normalized_header(value) for value in (
    "序号", "编号", "手机号", "联系电话", "电话", "身份证号", "邮箱", "班级", "专业", "年级",
)}
NAME_NEGATIVE_HEADERS = {_normalized_header(value) for value in ("班级名称", "专业名称", "课程名称", "学院名称")}


@dataclass(frozen=True)
class StudentRosterRow:
    row_no: int
    student_id: str
    name: str
    password: str


@dataclass(frozen=True)
class ParsedStudentRoster:
    rows: list[StudentRosterRow]
    sheet_name: str
    header_row: int
    student_id_column: str
    name_column: str
    password_column: str | None
    ignored_columns: int


def _cell_text(value, number_format: str = "") -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        digits = re.fullmatch(r"0+", (number_format or "").split(";", 1)[0].strip())
        return f"{value:0{len(digits.group(0))}d}" if digits else str(value)
    if isinstance(value, float):
        if value.is_integer():
            integer = int(value)
            digits = re.fullmatch(r"0+", (number_format or "").split(";", 1)[0].strip())
            return f"{integer:0{len(digits.group(0))}d}" if digits else str(integer)
        return format(value, ".15g")
    return str(value).strip()


def _read_xlsx(data: bytes) -> list[tuple[str, list[tuple[int, list[str]]]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请确认文件未损坏且未加密") from exc
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row_no, cells in enumerate(sheet.iter_rows(max_col=MAX_ROSTER_COLUMNS), start=1):
                if row_no > MAX_ROSTER_ROWS + HEADER_SCAN_ROWS:
                    break
                values = [_cell_text(cell.value, cell.number_format) for cell in cells]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append((row_no, values))
            if rows:
                sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def _read_xls(data: bytes) -> list[tuple[str, list[tuple[int, list[str]]]]]:
    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        raise ValueError("旧版 XLS 文件无法读取，请确认文件未损坏且未加密") from exc
    sheets = []
    try:
        for sheet in workbook.sheets():
            rows = []
            for row_index in range(min(sheet.nrows, MAX_ROSTER_ROWS + HEADER_SCAN_ROWS)):
                values = [_cell_text(sheet.cell_value(row_index, column)) for column in range(min(sheet.ncols, MAX_ROSTER_COLUMNS))]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append((row_index + 1, values))
            if rows:
                sheets.append((sheet.name, rows))
    finally:
        workbook.release_resources()
    return sheets


def _header_score(header: str, aliases: set[str], negatives: set[str] | None = None) -> int:
    normalized = _normalized_header(header)
    if not normalized:
        return 0
    if negatives and any(item == normalized or item in normalized for item in negatives):
        return -100
    if normalized in aliases:
        return 120
    if any(len(alias) >= 2 and (alias in normalized or normalized in alias) for alias in aliases):
        return 80
    return 0


def _column_values(rows: list[tuple[int, list[str]]], column: int) -> list[str]:
    return [values[column].strip() for _, values in rows if column < len(values) and values[column].strip()][:200]


def _id_data_score(values: list[str]) -> float:
    if len(values) < 2:
        return 0
    id_like = sum(bool(re.fullmatch(r"(?=.*\d)[A-Za-z0-9_-]{4,40}", value)) for value in values) / len(values)
    unique = len(set(values)) / len(values)
    digit_bearing = sum(any(character.isdigit() for character in value) for value in values) / len(values)
    sequential = False
    if all(value.isdigit() for value in values):
        numbers = [int(value) for value in values]
        sequential = len(numbers) >= 3 and numbers == list(range(numbers[0], numbers[0] + len(numbers)))
    return id_like * 60 + unique * 25 + digit_bearing * 15 - (10 if sequential else 0)


def _name_data_score(values: list[str]) -> float:
    if len(values) < 2:
        return 0
    def looks_like_name(value: str) -> bool:
        return bool(
            re.fullmatch(r"[\u3400-\u9fff·]{2,12}", value)
            or re.fullmatch(r"[A-Za-z][A-Za-z .·'-]{1,39}", value)
        )
    name_like = sum(looks_like_name(value) for value in values) / len(values)
    unique = len(set(values)) / len(values)
    return name_like * 80 + unique * 20


def _column_name(header: list[str], column: int) -> str:
    return header[column].strip() if column < len(header) and header[column].strip() else f"第 {column + 1} 列"


def _best_sheet_mapping(sheets):
    candidates = []
    for sheet_name, sheet_rows in sheets:
        for header_index in range(min(HEADER_SCAN_ROWS, max(0, len(sheet_rows) - 2))):
            header_row_no, header = sheet_rows[header_index]
            data_rows = sheet_rows[header_index + 1:]
            if sum(bool(value.strip()) for value in header) < 2 or len(data_rows) < 2:
                continue
            column_count = min(MAX_ROSTER_COLUMNS, max(len(header), *(len(values) for _, values in data_rows[:200])))
            id_scores = []
            name_scores = []
            for column in range(column_count):
                values = _column_values(data_rows, column)
                label = header[column] if column < len(header) else ""
                id_scores.append(_header_score(label, ID_ALIASES, ID_NEGATIVE_HEADERS) + _id_data_score(values))
                name_scores.append(_header_score(label, NAME_ALIASES, NAME_NEGATIVE_HEADERS) + _name_data_score(values))
            pairs = [
                (id_scores[id_column] + name_scores[name_column], id_column, name_column)
                for id_column in range(column_count)
                for name_column in range(column_count)
                if id_column != name_column and id_scores[id_column] >= 65 and name_scores[name_column] >= 65
            ]
            if not pairs:
                continue
            pairs.sort(reverse=True)
            score, id_column, name_column = pairs[0]
            ambiguous_columns = len(pairs) > 1 and score - pairs[1][0] < 12 and (id_column, name_column) != (pairs[1][1], pairs[1][2])
            password_scores = [
                (_header_score(header[column] if column < len(header) else "", PASSWORD_ALIASES), column)
                for column in range(column_count)
                if column not in {id_column, name_column}
            ]
            password_score, password_column = max(password_scores, default=(0, -1))
            candidates.append({
                "score": score,
                "sheet_name": sheet_name,
                "rows": data_rows,
                "header": header,
                "header_row": header_row_no,
                "id_column": id_column,
                "name_column": name_column,
                "password_column": password_column if password_score >= 80 else None,
                "ambiguous_columns": ambiguous_columns,
            })
    if not candidates:
        raise ValueError("无法识别学号和姓名列；请保留表头，并确保数据至少有两名学生")
    candidates.sort(key=lambda item: item["score"], reverse=True)
    best = candidates[0]
    if best["ambiguous_columns"]:
        raise ValueError("检测到多个可能的学号/姓名列，无法安全自动选择；请将对应表头改为“学号”和“姓名”后重试")
    if len(candidates) > 1:
        second = candidates[1]
        mapping_differs = (best["sheet_name"], best["header_row"], best["id_column"], best["name_column"]) != (
            second["sheet_name"], second["header_row"], second["id_column"], second["name_column"]
        )
        if mapping_differs and best["score"] - second["score"] < 12:
            raise ValueError("检测到多个可能的学号/姓名列，无法安全自动选择；请将对应表头改为“学号”和“姓名”后重试")
    return best


def parse_student_roster(upload, default_password: str = "") -> ParsedStudentRoster:
    filename = upload.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise ValueError("学生名单仅支持 Excel .xlsx 或 .xls 文件")
    data = upload.stream.read(MAX_ROSTER_BYTES + 1)
    upload.stream.seek(0)
    if not data:
        raise ValueError("上传的学生名单为空")
    if len(data) > MAX_ROSTER_BYTES:
        raise ValueError("学生名单不能超过 10 MB")
    if suffix == ".xlsx" and not data.startswith(b"PK"):
        raise ValueError("文件扩展名为 .xlsx，但内容不是有效的 Excel 工作簿")
    if suffix == ".xls" and not data.startswith(bytes.fromhex("D0CF11E0")):
        raise ValueError("文件扩展名为 .xls，但内容不是有效的旧版 Excel 工作簿")
    sheets = _read_xlsx(data) if suffix == ".xlsx" else _read_xls(data)
    if not sheets:
        raise ValueError("Excel 中没有可读取的数据工作表")
    mapping = _best_sheet_mapping(sheets)
    rows = []
    duplicates = {}
    invalid_rows = []
    for row_no, values in mapping["rows"]:
        student_id = values[mapping["id_column"]].strip() if mapping["id_column"] < len(values) else ""
        name = values[mapping["name_column"]].strip() if mapping["name_column"] < len(values) else ""
        if not student_id and not name:
            continue
        if not student_id or not name:
            invalid_rows.append(row_no)
            continue
        if not re.fullmatch(r"[^\s,]{2,64}", student_id):
            raise ValueError(f"第 {row_no} 行学号格式无效")
        if len(name) > 120:
            raise ValueError(f"第 {row_no} 行姓名超过 120 个字符")
        password = ""
        if mapping["password_column"] is not None and mapping["password_column"] < len(values):
            password = values[mapping["password_column"]].strip()
        password = password or default_password.strip()
        if student_id in duplicates:
            raise ValueError(f"学号 {student_id} 在第 {duplicates[student_id]} 行和第 {row_no} 行重复")
        duplicates[student_id] = row_no
        rows.append(StudentRosterRow(row_no=row_no, student_id=student_id, name=name, password=password))
        if len(rows) > MAX_ROSTER_ROWS:
            raise ValueError(f"单次最多导入 {MAX_ROSTER_ROWS} 名学生")
    if invalid_rows:
        preview = "、".join(str(value) for value in invalid_rows[:5])
        raise ValueError(f"第 {preview} 行的学号或姓名为空，请补全后重试")
    if not rows:
        raise ValueError("Excel 中没有识别到可导入的学生记录")
    header = mapping["header"]
    selected = {mapping["id_column"], mapping["name_column"]}
    if mapping["password_column"] is not None:
        selected.add(mapping["password_column"])
    ignored = sum(bool(value.strip()) for column, value in enumerate(header) if column not in selected)
    return ParsedStudentRoster(
        rows=rows,
        sheet_name=mapping["sheet_name"],
        header_row=mapping["header_row"],
        student_id_column=_column_name(header, mapping["id_column"]),
        name_column=_column_name(header, mapping["name_column"]),
        password_column=_column_name(header, mapping["password_column"]) if mapping["password_column"] is not None else None,
        ignored_columns=ignored,
    )
