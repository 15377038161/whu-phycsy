from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select

from .models import (
    Course,
    CourseExperiment,
    Enrollment,
    Experiment,
    ExperimentVersion,
    QuizAssignment,
    Review,
    SubmissionRevision,
    User,
)


WARNING_DAYS = 7


def _aware(value: datetime | None, field: str = "时间") -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        raise ValueError(f"{field}必须包含时区")
    return value.astimezone(timezone.utc)


def _stored(value: datetime | None) -> datetime | None:
    """SQLite may return persisted UTC timestamps without tzinfo."""
    if value is None:
        return None
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _quiz_score(assignment: QuizAssignment | None) -> int | None:
    if not assignment or not assignment.completed_at:
        return None
    questions = assignment.questions or []
    answers = assignment.answers or {}
    if not questions:
        return None
    correct = sum(answers.get(question.get("id")) == question.get("answer") for question in questions)
    return round(correct * 100 / len(questions))


def _average(values: list[int]) -> float | None:
    return round(sum(values) / len(values), 1) if values else None


def _empty(courses: list[Course], now: datetime, course: Course | None = None) -> dict:
    summary = {
        "student_count": 0,
        "visible_student_count": 0,
        "course_experiment_count": 0,
        "task_count": 0,
        "assignment_count": 0,
        "completed": 0,
        "in_progress": 0,
        "not_started": 0,
        "completed_task_cells": 0,
        "in_progress_task_cells": 0,
        "not_started_task_cells": 0,
        "completion_rate": 0.0,
        "in_progress_rate": 0.0,
        "not_started_rate": 0.0,
        "completions_in_range": 0,
        "quiz_average": None,
        "teacher_score_average": None,
        "warning_count": 0,
    }
    return {
        "scope": {
            "teacher_id": course.teacher_id if course else None,
            "course_id": course.id if course else None,
            "course_name": course.name if course else None,
        },
        "courses": [{"id": item.id, "name": item.name} for item in courses],
        "course": {"id": course.id, "name": course.name, "teacher_id": course.teacher_id} if course else None,
        "generated_at": now.isoformat(),
        "filters": {},
        "summary": summary,
        "tasks": [],
        "students": [],
        "experiment_stats": [],
        "slowest_students": [],
        "student_lists": {"selected_task_id": None, "completed": [], "incomplete": []},
        "trends": [
            {"date": (now.date() - timedelta(days=i)).isoformat(), "prestudy": 0, "report": 0}
            for i in range(6, -1, -1)
        ],
        "trend": [
            {"date": (now.date() - timedelta(days=i)).isoformat(), "completed": 0}
            for i in range(6, -1, -1)
        ],
        "alerts": [{"kind": "empty", "title": "当前筛选范围暂无学习活动"}],
    }


def get_teacher_analytics(
    session,
    teacher_id: str,
    *,
    course_id: str | None = None,
    experiment_id: str | None = None,
    task_type: str | None = None,
    status: str | None = None,
    submitted_from: datetime | None = None,
    submitted_to: datetime | None = None,
    now: datetime | None = None,
) -> dict:
    now = _aware(now, "当前时间") if now else datetime.now(timezone.utc)
    start, end = _aware(submitted_from, "开始时间"), _aware(submitted_to, "结束时间")
    if start and end and start > end:
        raise ValueError("开始时间不能晚于结束时间")
    if task_type not in {None, "prestudy", "report"} or status not in {
        None,
        "completed",
        "in_progress",
        "not_started",
    }:
        raise ValueError("筛选条件无效")

    def within_range(value: datetime | None) -> bool:
        stored = _stored(value)
        return bool(stored and (start is None or stored >= start) and (end is None or stored <= end))

    courses = session.scalars(
        select(Course).where(Course.teacher_id == teacher_id).order_by(Course.created_at, Course.id)
    ).all()
    if not courses:
        result = _empty([], now)
        result["scope"]["teacher_id"] = teacher_id
        return result
    owned = {item.id: item for item in courses}
    if course_id and course_id not in owned:
        raise ValueError("课程不属于当前教师")
    course = owned.get(course_id) if course_id else courses[0]

    roster = session.scalars(
        select(User)
        .join(Enrollment, Enrollment.student_id == User.id)
        .where(Enrollment.course_id == course.id, User.role == "student", User.active.is_(True))
        .order_by(User.id)
    ).all()
    experiment_rows = session.execute(
        select(CourseExperiment, Experiment)
        .join(Experiment, Experiment.id == CourseExperiment.experiment_id)
        .where(
            CourseExperiment.course_id == course.id,
            CourseExperiment.active.is_(True),
            Experiment.retired.is_(False),
            *([Experiment.id == experiment_id] if experiment_id else []),
        )
        .order_by(CourseExperiment.position, Experiment.order_no)
    ).all()
    if not roster or not experiment_rows:
        result = _empty(courses, now, course)
        result["summary"]["student_count"] = len(roster)
        result["summary"]["visible_student_count"] = len(roster)
        result["summary"]["course_experiment_count"] = len(experiment_rows)
        return result

    experiment_ids = [experiment.id for _, experiment in experiment_rows]
    version_rows = session.execute(
        select(ExperimentVersion.id, ExperimentVersion.experiment_id).where(
            ExperimentVersion.experiment_id.in_(experiment_ids)
        )
    ).all()
    version_exp = dict(version_rows)
    version_ids = list(version_exp)
    student_ids = [student.id for student in roster]
    quizzes = session.scalars(
        select(QuizAssignment).where(
            QuizAssignment.course_id == course.id,
            QuizAssignment.student_id.in_(student_ids),
            QuizAssignment.experiment_version_id.in_(version_ids),
        )
    ).all()
    revisions = session.scalars(
        select(SubmissionRevision)
        .where(
            SubmissionRevision.course_id == course.id,
            SubmissionRevision.student_id.in_(student_ids),
            SubmissionRevision.experiment_version_id.in_(version_ids),
        )
        .order_by(SubmissionRevision.submitted_at.desc(), SubmissionRevision.revision_no.desc())
    ).all()
    reviews = session.scalars(
        select(Review)
        .where(
            Review.submission_id.in_([revision.id for revision in revisions]),
            Review.superseded.is_(False),
        )
        .order_by(Review.created_at.desc())
    ).all() if revisions else []

    quiz_map: dict[tuple[str, str], QuizAssignment] = {}
    for quiz in quizzes:
        key = (quiz.student_id, version_exp[quiz.experiment_version_id])
        if key not in quiz_map or _stored(quiz.frozen_at) > _stored(quiz_map[key].frozen_at):
            quiz_map[key] = quiz
    revision_map: dict[tuple[str, str], SubmissionRevision] = {}
    for revision in revisions:
        revision_map.setdefault((revision.student_id, version_exp[revision.experiment_version_id]), revision)
    review_map: dict[str, Review] = {}
    for review in reviews:
        review_map.setdefault(review.submission_id, review)

    definitions = [
        (f"{experiment.id}:{kind}", experiment, kind)
        for _, experiment in experiment_rows
        for kind in ("prestudy", "report")
        if task_type in {None, kind}
    ]
    matrix = {student.id: {} for student in roster}
    trend_counts = Counter()
    completions_in_range = 0
    for student in roster:
        for task_id, experiment, kind in definitions:
            quiz = quiz_map.get((student.id, experiment.id))
            revision = revision_map.get((student.id, experiment.id))
            if kind == "prestudy":
                completed_at = _stored(quiz.completed_at) if quiz and quiz.completed_at else None
                state = "completed" if completed_at else ("in_progress" if quiz else "not_started")
            else:
                completed_at = _stored(revision.submitted_at) if revision and revision.status == "submitted" else None
                state = "completed" if completed_at else (
                    "in_progress" if revision or (quiz and quiz.completed_at) else "not_started"
                )
            in_range = bool(completed_at and (start is None or completed_at >= start) and (end is None or completed_at <= end))
            if completed_at and (start or end) and not in_range:
                state, completed_at = "not_started", None
            if completed_at:
                trend_counts[(completed_at.date(), kind)] += 1
                completions_in_range += 1
            matrix[student.id][task_id] = {
                "status": state,
                "completed_at": completed_at.isoformat() if completed_at else None,
            }

    tasks = []
    for task_id, experiment, kind in definitions:
        groups = {key: [] for key in ("completed", "in_progress", "not_started")}
        for student in roster:
            entry = matrix[student.id][task_id]
            groups[entry["status"]].append(
                {
                    "id": student.id,
                    "student_id": student.id,
                    "name": student.name,
                    "student_name": student.name,
                    "status": entry["status"],
                    "completed_at": entry["completed_at"],
                }
            )
        if status and not groups[status]:
            continue
        counts = {key: len(value) for key, value in groups.items()}
        total = len(roster)
        tasks.append(
            {
                "id": task_id,
                "task_id": task_id,
                "experiment_id": experiment.id,
                "experiment_code": experiment.code,
                "experiment": experiment.title,
                "task_title": f"{experiment.title}·{'预习' if kind == 'prestudy' else '报告'}",
                "type": kind,
                "task_type": kind,
                "type_label": "预习" if kind == "prestudy" else "报告",
                "total": total,
                **counts,
                "counts": counts,
                "completed_students": groups["completed"],
                "in_progress_students": groups["in_progress"],
                "not_started_students": groups["not_started"],
                "completion_rate": round(counts["completed"] * 100 / total, 1),
                "completions_in_range": len(groups["completed"]),
            }
        )

    visible_ids = {task["id"] for task in tasks}
    students = []
    all_quiz_scores: list[int] = []
    all_teacher_scores: list[int] = []
    for student in roster:
        entries = [
            {"task_id": task_id, **entry}
            for task_id, entry in matrix[student.id].items()
            if task_id in visible_ids
        ]
        if status and not any(entry["status"] == status for entry in entries):
            continue
        counts = Counter(entry["status"] for entry in entries)
        overall = "completed" if entries and counts["completed"] == len(entries) else (
            "in_progress" if counts["completed"] or counts["in_progress"] else "not_started"
        )
        student_quiz_scores = [
            score
            for experiment_id_item in experiment_ids
            if task_type in {None, "prestudy"}
            and (quiz_item := quiz_map.get((student.id, experiment_id_item)))
            and within_range(quiz_item.completed_at)
            and (score := _quiz_score(quiz_item)) is not None
        ]
        current_revisions = [
            revision_map.get((student.id, experiment_id_item)) for experiment_id_item in experiment_ids
        ]
        student_teacher_scores = [
            review_map[revision.id].private_score
            for revision in current_revisions
            if task_type in {None, "report"}
            and revision
            and revision.status == "submitted"
            and within_range(revision.submitted_at)
            and revision.id in review_map
        ]
        activity_values = [
            _stored(value)
            for value in [
                *[quiz.frozen_at for (student_id, _), quiz in quiz_map.items() if student_id == student.id],
                *[revision.submitted_at for revision in current_revisions if revision],
            ]
            if value
        ]
        latest_activity = max(activity_values) if activity_values else None
        inactive_days = max(0, (now - latest_activity).days) if latest_activity else None
        warning = overall != "completed" and (inactive_days is None or inactive_days >= WARNING_DAYS)
        warning_reason = (
            "尚未开始任何任务" if warning and inactive_days is None else
            f"未完成且已 {inactive_days} 天无新活动" if warning else ""
        )
        completion_rate = round(counts["completed"] * 100 / len(entries), 1) if entries else 0.0
        all_quiz_scores.extend(student_quiz_scores)
        all_teacher_scores.extend(student_teacher_scores)
        students.append(
            {
                "id": student.id,
                "student_id": student.id,
                "username": student.username,
                "name": student.name,
                "student_name": student.name,
                "status": overall,
                "tasks": entries,
                "completed": counts["completed"],
                "in_progress": counts["in_progress"],
                "not_started": counts["not_started"],
                "total": len(entries),
                "completion_rate": completion_rate,
                "quiz_average": _average(student_quiz_scores),
                "teacher_score_average": _average(student_teacher_scores),
                "latest_activity_at": latest_activity.isoformat() if latest_activity else None,
                "inactive_days": inactive_days,
                "warning": warning,
                "warning_reason": warning_reason,
            }
        )

    experiment_stats = []
    task_by_id = {task["id"]: task for task in tasks}
    for _, experiment in experiment_rows:
        prestudy = task_by_id.get(f"{experiment.id}:prestudy")
        report = task_by_id.get(f"{experiment.id}:report")
        quiz_scores = [
            score
            for student in roster
            if task_type in {None, "prestudy"}
            and (quiz_item := quiz_map.get((student.id, experiment.id)))
            and within_range(quiz_item.completed_at)
            and (score := _quiz_score(quiz_item)) is not None
        ]
        teacher_scores = []
        submitted_times = []
        for student in roster:
            revision = revision_map.get((student.id, experiment.id))
            if revision and revision.status == "submitted" and within_range(revision.submitted_at):
                if task_type in {None, "report"}:
                    submitted_times.append(_stored(revision.submitted_at))
                if task_type in {None, "report"} and revision.id in review_map:
                    teacher_scores.append(review_map[revision.id].private_score)
        experiment_stats.append(
            {
                "experiment_id": experiment.id,
                "experiment_code": experiment.code,
                "experiment": experiment.title,
                "prestudy_rate": prestudy["completion_rate"] if prestudy else None,
                "report_rate": report["completion_rate"] if report else None,
                "quiz_average": _average(quiz_scores),
                "teacher_score_average": _average(teacher_scores),
                "pending_reviews": len(submitted_times) - len(teacher_scores),
                "latest_submission_at": max(submitted_times).isoformat() if submitted_times else None,
            }
        )

    cell_counts = Counter(
        entry["status"]
        for student in roster
        for task_id, entry in matrix[student.id].items()
        if task_id in visible_ids
    )
    assignments = sum(cell_counts.values())
    selected = tasks[0] if tasks else None
    warning_students = [student for student in students if student["warning"]]
    slowest_students = sorted(
        students,
        key=lambda item: (
            item["completion_rate"],
            item["latest_activity_at"] is not None,
            item["latest_activity_at"] or "",
            item["id"],
        ),
    )[:3]
    trends = [
        {
            "date": (now.date() - timedelta(days=i)).isoformat(),
            "prestudy": trend_counts[(now.date() - timedelta(days=i), "prestudy")],
            "report": trend_counts[(now.date() - timedelta(days=i), "report")],
        }
        for i in range(6, -1, -1)
    ]
    summary = {
        "student_count": len(roster),
        "visible_student_count": len(students),
        "course_experiment_count": len(experiment_rows),
        "task_count": len(tasks),
        "assignment_count": assignments,
        "completed": cell_counts["completed"],
        "in_progress": cell_counts["in_progress"],
        "not_started": cell_counts["not_started"],
        "completed_task_cells": cell_counts["completed"],
        "in_progress_task_cells": cell_counts["in_progress"],
        "not_started_task_cells": cell_counts["not_started"],
        "completion_rate": round(cell_counts["completed"] * 100 / assignments, 1) if assignments else 0.0,
        "in_progress_rate": round(cell_counts["in_progress"] * 100 / assignments, 1) if assignments else 0.0,
        "not_started_rate": round(cell_counts["not_started"] * 100 / assignments, 1) if assignments else 0.0,
        "completions_in_range": completions_in_range,
        "quiz_average": _average(all_quiz_scores),
        "teacher_score_average": _average(all_teacher_scores),
        "warning_count": len(warning_students),
    }
    return {
        "scope": {"teacher_id": teacher_id, "course_id": course.id, "course_name": course.name},
        "courses": [{"id": item.id, "name": item.name} for item in courses],
        "course": {"id": course.id, "name": course.name, "teacher_id": teacher_id},
        "generated_at": now.isoformat(),
        "filters": {
            "experiment_id": experiment_id,
            "task_type": task_type,
            "status": status,
            "submitted_from": start.isoformat() if start else None,
            "submitted_to": end.isoformat() if end else None,
        },
        "summary": summary,
        "tasks": tasks,
        "students": students,
        "experiment_stats": experiment_stats,
        "slowest_students": slowest_students,
        "student_lists": {
            "selected_task_id": selected["id"] if selected else None,
            "completed": selected["completed_students"] if selected else [],
            "incomplete": selected["in_progress_students"] + selected["not_started_students"] if selected else [],
        },
        "trends": trends,
        "trend": [{"date": point["date"], "completed": point["prestudy"] + point["report"]} for point in trends],
        "alerts": [
            {
                "kind": "student_warning",
                "student_id": student["id"],
                "title": student["name"],
                "detail": student["warning_reason"],
            }
            for student in warning_students
        ][:10] or [{"kind": "clear", "title": "当前没有停滞预警"}],
    }


def build_teacher_analytics_workbook(payload: dict) -> BytesIO:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "班级概览"
    overview.append(["班级", "学生数", "整体完成率", "预习平均分", "教师评分均分", "预警人数", "生成时间"])
    summary = payload["summary"]
    overview.append([
        payload.get("course", {}).get("name", ""),
        summary["student_count"],
        summary["completion_rate"] / 100,
        summary["quiz_average"],
        summary["teacher_score_average"],
        summary["warning_count"],
        payload["generated_at"],
    ])
    overview["C2"].number_format = "0.0%"

    detail = workbook.create_sheet("学生学情明细")
    detail.append(["学号", "姓名", "状态", "完成任务", "任务总数", "完成率", "预习平均分", "教师评分均分", "最近活动", "预警"])
    status_labels = {"completed": "已完成", "in_progress": "进行中", "not_started": "未开始"}
    for student in payload["students"]:
        detail.append([
            student["username"],
            student["name"],
            status_labels.get(student["status"], student["status"]),
            student["completed"],
            student["total"],
            student["completion_rate"] / 100,
            student["quiz_average"],
            student["teacher_score_average"],
            student["latest_activity_at"] or "—",
            student["warning_reason"] or "—",
        ])
        detail.cell(detail.max_row, 6).number_format = "0.0%"

    experiments = workbook.create_sheet("实验任务完成率")
    experiments.append(["实验", "预习完成率", "报告完成率", "预习平均分", "教师评分均分", "待审核", "最近提交"])
    for item in payload["experiment_stats"]:
        experiments.append([
            item["experiment"],
            item["prestudy_rate"] / 100 if item["prestudy_rate"] is not None else None,
            item["report_rate"] / 100 if item["report_rate"] is not None else None,
            item["quiz_average"],
            item["teacher_score_average"],
            item["pending_reviews"],
            item["latest_submission_at"] or "—",
        ])
        experiments.cell(experiments.max_row, 2).number_format = "0.0%"
        experiments.cell(experiments.max_row, 3).number_format = "0.0%"

    header_fill = PatternFill("solid", fgColor="0B3558")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = min(36, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _without_scores(value):
    if isinstance(value, dict):
        return {
            key: _without_scores(item)
            for key, item in value.items()
            if key not in {"teacher_score_average", "quiz_average"}
        }
    if isinstance(value, list):
        return [_without_scores(item) for item in value]
    return value


def build_teacher_ai_snapshot(source, teacher_id: str | None = None, **filters) -> dict:
    payload = get_teacher_analytics(source, teacher_id, **filters) if teacher_id is not None else source
    snapshot = {
        key: payload.get(key)
        for key in ("scope", "generated_at", "filters", "summary", "tasks", "students", "trends", "alerts")
    }
    return _without_scores(snapshot)
