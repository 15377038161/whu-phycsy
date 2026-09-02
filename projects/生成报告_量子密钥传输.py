# -*- coding: utf-8 -*-
"""
生成量子密钥传输实验报告 Word 文档
严格对齐示例 PDF 格式：
- A4 纸张 (210mm x 297mm)
- DengXian（等线）字体
- 标题 18pt Bold，章节标题 12pt Bold，正文 10.5pt
- 中文编号（一、二、三...）
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

doc = Document()

# ============ 页面设置 ============
section = doc.sections[0]
section.page_width = Cm(21.0)    # A4 宽度
section.page_height = Cm(29.7)   # A4 高度
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.17)
section.right_margin = Cm(3.17)

# ============ 默认样式 ============
style = doc.styles['Normal']
font = style.font
font.name = 'DengXian'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0, 0, 0)
# 设置中文字体
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'DengXian')
# 行距
pf = style.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.25
pf.space_before = Pt(0)
pf.space_after = Pt(0)

def set_font(run, font_name='DengXian', size=10.5, bold=False, color=(0,0,0)):
    """设置字体（含中文）"""
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def add_title(text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    """添加居中标题"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, 'DengXian', size, bold)
    return p

def add_section_heading(text, size=12, bold=True):
    """添加章节标题（一、二、三...）"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, 'DengXian', size, bold)
    return p

def add_body(text, indent=False):
    """添加正文段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.25
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)  # 首行缩进2字符
    run = p.add_run(text)
    set_font(run, 'DengXian', 10.5, False)
    return p

def add_table_with_data(headers, rows, col_widths=None):
    """添加带数据的表格"""
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 设置表格字体
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_font(run, 'DengXian', 9, True)

    # 表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        set_font(run, 'DengXian', 9, True)

    # 数据行
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_data in enumerate(row_data):
            cell = table.rows[row_idx+1].cells[col_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(cell_data))
            set_font(run, 'DengXian', 9, False)

    # 设置列宽
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)

    return table

# ============ 文档内容 ============

# 标题
add_title('量子密钥传输实验', 18, True)

# 作者和日期
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run('范文轩 2022300002071')
set_font(run, 'DengXian', 10.5, False)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
run = p.add_run('2025年10月15日')
set_font(run, 'DengXian', 10.5, False)

# 一、实验目的
add_section_heading('一、实验目的')
add_body('1. 了解量子密码的基本流程和安全性；')
add_body('2. 理解基于BB84协议实现技术细节；')
add_body('3. 了解QKD的研究进展;')
add_body('4. 了解量子密钥分发系统的实现方式和实验控制流程;')
add_body('5. 进行量子密钥传输实验。')

# 二、主要实验仪器
add_section_heading('二、主要实验仪器')
add_body('红外QKD教学机主要包括制备单个光子的单光子源，量子态的调制，实现光传输的光纤信道，实现单光子探测的单光子探测器，实现密钥分发收发同步的同步系统，以及整个系统的控制软件和密钥生成软件。')

# 三、实验原理
add_section_heading('三、实验原理')
add_body('量子密钥分发实验中，协议实施的流程大致描述如下：')
add_body('发送方Alice制备一系列的光子发送给接收方Bob，每个光子的偏振态独立随机地从水平偏振态|→⟩、竖直偏振态|↑⟩、右斜45度偏振态|↗⟩和左斜45度偏振态|↖⟩四个偏振态中选取，如果Alice发送光子的偏振态为水平偏振态|→⟩或者竖直偏振态|↑⟩，则称Alice选择+基制备光子，如果Alice发送光子的偏振态是右斜45度偏振态|↗⟩或者左斜45度偏振态|↖⟩，则称Alice选择×基制备光子。')
add_body('接收方Bob与Alice完全独立地随机选取+基和X基测量Alice发送过来光子的偏振态，并记录下测量到光子的位置信息。')
add_body('Alice和Bob对基，即双方仅保留基相同（Alice制备基和Bob测量基）并且Bob测量到光子位置的光子偏振态信息，双方基不同时则直接抛弃相关信息。')
add_body('Alice和Bob将保留的光子偏振态信息转换成相应的密钥比特信息，即对基后保留的光子偏振态按水平偏振态|→⟩和右斜45度偏振态|↗⟩转换为比特"0"，竖直偏振态|↑⟩和左斜45度偏振态|↖⟩转换为比特"1"。')
add_body('Alice和Bob通过经典公开信道对上一步中获得的密钥比特进行处理，其过程主要分成纠错和保密放大来进行，纠错就是使得密钥比特一致，而保密放大（Privacy Amplification）就是将可能泄漏给窃听者的信息剔除掉。')
add_body('量子密钥分发系统中，用两位的bit编码表示光子信息，其中个位bit代表基矢信息，十位bit代表密钥信息；例如：Alice端水平偏振编码为00，垂直偏振编码为10，右斜45度偏振编码为01，左斜45度偏振编码为11；相应的Bob端四路探测器探测到信号，分别也是按照上述编码方式进行编码。')
add_body('密钥分发的过程中，光子传输探测后，会得到一系列的这种两位编码的信息数据，如何从这些数据中提取出有用信息，需要原始数据（raw）经过对基（sift）、纠错（reconcile）、保密放大等过程，最后得到安全密钥（key），以保证密钥的安全性。当然，考虑到是否存在窃听，需要对系统的每次传输过程进行误码估计。以保证此次的传输数据有效。')
add_body('Alice和Bob两端传输探测完成后会得到一系列两位编码的信息数据，首先需要对两端的数据进行对基，再对对基的数据进行比对，计算出系统的误码率。当误码率低于理论安全界限11%时，本次传输有效，继续进行后续处理过程。系统对基的过程实例如下图所示。')

# 图1说明
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)
run = p.add_run('图1 密钥分发流程示意图')
set_font(run, 'DengXian', 10.5, False)

# 四、实验内容与步骤
add_section_heading('四、实验内容与步骤')
add_body('1、发送端和接收端电路控制模块通电，启动QKD软件；')
add_body('2、点击QKD的发送端控制界面（Alice-Bob），勾选中间密钥输出功能；点击QKD的运行按钮；待系统运行5-10秒后，点击停止按钮。此时，记录下系统的平均误码率；')
add_body('3、在桌面的自由空间偏振文件夹中查看中间密钥输出的数据，输出的数据主要包括原始数据（raw），对基后数据（sift），纠错后数据（reconcile），最终安全密钥数据（key），对应的发送端分别为：transmitter-raw、transmitter-sift、transmitter-reconcile、transmitter-key，接收端分别为：receiver-raw、receiver-sift、receiver-reconcile、receiver-key；')
add_body('4、打开桌面上的量芯对比工具，选择文本比较，进入界面，点击最上方的会话选项，比较文件；')
add_body('5、使用文件对比工具分别打开transmitter-sift和receiver-sift，然后从中选择部分数据（例如中间位置10行，若总行数为1000，则提取数据比例为1%）进行两端的误码估计，数出错误的个数（黄色显示）并计算误码率。（例如总共100个二进制数据，两端有5个错误位，则误码率为5%）')
add_body('6、根据实验原理的系统对基过程，依照下表的模拟数据，进行填空，以了解和学习密钥分发过程的数据处理。')
add_body('7、关闭系统硬件和软件。')

# 五、实验数据
add_section_heading('五、实验数据')

# 表1
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(4)
run = p.add_run('表1 QKD系统误码率数据')
set_font(run, 'DengXian', 10.5, True)

table1_headers = ['序号', '名称', '数据', '单位']
table1_rows = [
    ['1', 'QKD系统软件统计误码率', '5.457', '%'],
    ['2', 'QKD系统误码估计采样率', '20', '%'],
    ['3', '手动提取对基后数据比例', '1.68', '%'],
    ['4', '对基后数据误码率', '10.9', '%'],
]
add_table_with_data(table1_headers, table1_rows, [1.5, 6, 2, 1.5])

# 表2
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(4)
run = p.add_run('表2 密钥分发过程数据处理分析')
set_font(run, 'DengXian', 10.5, True)

positions = list(range(1, 17))
pos_headers = ['位置'] + [str(i) for i in positions]

a_encoding = ['00','01','10','11','11','01','00','10','10','11','01','00','00','10','11','01']
a_basis    = ['+','×','+','×','×','×','+','+','+','×','×','+','+','+','×','×']
a_key      = ['0','0','1','1','1','0','0','1','1','1','0','0','0','1','1','0']
polar      = ['→','↗','↑','↖','↖','↗','→','↑','↑','↖','↗','→','→','↑','↖','↗']
bob_basis  = ['×','×','+','+','+','×','+','×','×','+','×','+','×','+','+','×']
b_detect   = ['0','0','1','1','0','0','0','0','1','1','0','0','0','1','1','0']
basis_match= ['','√','√','','','√','√','','','','√','√','','√','','√']
a_sift     = ['','0','1','','','0','0','','','','0','0','','1','','0']
b_sift     = ['','0','1','','','0','0','','','','0','0','','1','','0']
error_bit  = ['','0','0','','','0','0','','','','0','0','','0','','0']

table2_rows = [
    ['A编码'] + a_encoding,
    ['A选基'] + a_basis,
    ['A密钥'] + a_key,
    ['光偏振态'] + polar,
    ['Bob选基'] + bob_basis,
    ['B探测'] + b_detect,
    ['对基'] + basis_match,
    ['A-Sift-key'] + a_sift,
    ['B-Sift-key'] + b_sift,
    ['Error位'] + error_bit,
]
add_table_with_data(pos_headers, table2_rows)

# 六、思考题
add_section_heading('六、思考题')

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p.paragraph_format.line_spacing = 1.25
run = p.add_run('a) 误码率的误差来源有哪些？')
set_font(run, 'DengXian', 10.5, True)

add_body('在量子密钥分发（QKD）实验中，误码率的误差来源包括以下几个方面：')
add_body('测量基不理想：偏振分束器等元件校准不准，造成测量基非完全正交。')
add_body('光学误差：在光纤传输中，由于应力、温度变化等因素，信道产生随机偏振旋转，改变光子的偏振态。')
add_body('光源不完美：实际使用弱相干态光源时，脉冲中存在多光子成分。')

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p.paragraph_format.line_spacing = 1.25
run = p.add_run('b) 对于随机性分析来说，采集数据的多少对随机性有什么影响？')
set_font(run, 'DengXian', 10.5, True)

add_body('在随机性分析中，采集数据的数量对随机性评估的准确性、可靠性和深度有决定性的影响。数据量不足时，任何关于随机性的结论都是脆弱、不可靠甚至完全错误的。数据量越大，分析越可信。大样本能揭示微小的、系统性的非随机性，比如极其微弱的相关性、频率偏差或周期性。小样本则无法检测到这些细微缺陷。小样本只能分析非常短的相关性（例如，相邻比特之间的关系）。大样本才能探测长程相关性、周期性或复杂的模式，这些可能在数万甚至数百万比特之后才显现。')

# 附实验数据记录单
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(4)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('附实验数据记录单')
set_font(run, 'DengXian', 12, True)

# 保存
output_path = '实例实验报告/量子密钥传输实验报告.docx'
doc.save(output_path)
print(f'已生成: {output_path}')
