from __future__ import annotations

import json
import io
from datetime import datetime, timedelta, timezone
from time import perf_counter

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from platform_app import create_app
from platform_app.db import db_session, transaction
from platform_app.models import (
    Course,
    CourseExperiment,
    Enrollment,
    Experiment,
    ExperimentVersion,
    QuizAssignment,
    Review,
    SubmissionRevision,
    User,
)
from platform_app.teacher_analytics import build_teacher_ai_snapshot, build_teacher_analytics_workbook, get_teacher_analytics


@pytest.fixture()
def analytics_app(tmp_path, monkeypatch):
    monkeypatch.delenv("AI_API_KEY", raising=False)
    app = create_app(
        {
            "TESTING": True,
            "APP_ENV": "testing",
            "DATABASE_URL": f"sqlite:///{tmp_path / 'analytics.db'}",
            "DATA_DIR": str(tmp_path / "files"),
            "SECRET_KEY": "teacher-analytics-test",
        }
    )
    yield app


def _revision(
    *,
    request_id: str,
    student_id: str,
    course_id: str,
    version_id: str,
    revision_no: int,
    status: str,
    submitted_at: datetime,
    withdrawn_at: datetime | None = None,
    secret: str = "report-body-must-stay-private",
) -> SubmissionRevision:
    return SubmissionRevision(
        request_id=request_id,
        student_id=student_id,
        course_id=course_id,
        experiment_version_id=version_id,
        revision_no=revision_no,
        status=status,
        payload={"abstract": secret, "file_hashes": ["attachment-hash-must-stay-private"]},
        result_value=1.0,
        relative_error=0.0,
        deterministic_score=100,
        passed=True,
        fingerprint=f"fp-{request_id}",
        submitted_at=submitted_at,
        withdrawn_at=withdrawn_at,
    )


@pytest.fixture()
def analytics_data(analytics_app):
    now = datetime(2026, 7, 23, 8, 0, tzinfo=timezone.utc)
    with analytics_app.app_context(), transaction() as session:
        session.add_all(
            [
                User(id="TA", username="teacher-a", name="甲教师", role="teacher", password_hash="x"),
                User(id="TB", username="teacher-b", name="乙教师", role="teacher", password_hash="x"),
                User(id="TC", username="teacher-c", name="无课程教师", role="teacher", password_hash="x"),
                User(id="A1", username="2026001", name="甲一", role="student", password_hash="x"),
                User(id="A2", username="2026002", name="甲二", role="student", password_hash="x"),
                User(id="A3", username="2026003", name="零提交学生", role="student", password_hash="x"),
                User(id="B1", username="2026999", name="乙班学生", role="student", password_hash="x"),
            ]
        )
        course_a = Course(
            id="CA",
            teacher_id="TA",
            name="甲教师量子实验班",
            created_at=now - timedelta(days=100),
        )
        course_a_empty = Course(
            id="CA-EMPTY",
            teacher_id="TA",
            name="甲教师空班",
            created_at=now - timedelta(days=50),
        )
        course_b = Course(
            id="CB",
            teacher_id="TB",
            name="乙教师量子实验班",
            created_at=now - timedelta(days=90),
        )
        session.add_all([course_a, course_a_empty, course_b])
        session.add_all(
            [
                Enrollment(course_id="CA", student_id="A1"),
                Enrollment(course_id="CA", student_id="A2"),
                Enrollment(course_id="CA", student_id="A3"),
                Enrollment(course_id="CB", student_id="B1"),
            ]
        )

        experiments = []
        versions = []
        for position, (code, title) in enumerate(
            (("ANA1", "动态实验一"), ("ANA2", "动态实验二"), ("ANA3", "动态实验三")),
            start=1,
        ):
            experiment = Experiment(code=code, title=title, order_no=100 + position)
            session.add(experiment)
            session.flush()
            version = ExperimentVersion(
                experiment_id=experiment.id,
                version_no=1,
                status="published",
                definition={},
                published_at=now - timedelta(days=120),
            )
            session.add(version)
            session.flush()
            experiments.append(experiment)
            versions.append(version)
            session.add(CourseExperiment(course_id="CA", experiment_id=experiment.id, position=position))
        session.add(CourseExperiment(course_id="CB", experiment_id=experiments[0].id, position=1))

        session.add_all(
            [
                QuizAssignment(
                    course_id="CA",
                    student_id="A1",
                    experiment_version_id=versions[0].id,
                    questions=[],
                    answers={},
                    signature="a1-e1",
                    frozen_at=now - timedelta(days=3),
                    completed_at=now - timedelta(days=2),
                ),
                QuizAssignment(
                    course_id="CA",
                    student_id="A2",
                    experiment_version_id=versions[0].id,
                    questions=[],
                    answers={},
                    signature="a2-e1",
                    frozen_at=now - timedelta(days=1),
                    completed_at=None,
                ),
                QuizAssignment(
                    course_id="CA",
                    student_id="A2",
                    experiment_version_id=versions[1].id,
                    questions=[],
                    answers={},
                    signature="a2-e2",
                    frozen_at=now - timedelta(hours=6),
                    completed_at=now - timedelta(hours=3),
                ),
                QuizAssignment(
                    course_id="CB",
                    student_id="B1",
                    experiment_version_id=versions[0].id,
                    questions=[],
                    answers={},
                    signature="b1-e1",
                    frozen_at=now - timedelta(hours=2),
                    completed_at=now - timedelta(hours=1),
                ),
            ]
        )

        a1_report = _revision(
            request_id="a1-e1-r1",
            student_id="A1",
            course_id="CA",
            version_id=versions[0].id,
            revision_no=1,
            status="submitted",
            submitted_at=now - timedelta(days=1),
        )
        a2_withdrawn = _revision(
            request_id="a2-e1-r1",
            student_id="A2",
            course_id="CA",
            version_id=versions[0].id,
            revision_no=1,
            status="withdrawn",
            submitted_at=now - timedelta(hours=18),
            withdrawn_at=now - timedelta(hours=12),
        )
        a2_old = _revision(
            request_id="a2-e2-r1",
            student_id="A2",
            course_id="CA",
            version_id=versions[1].id,
            revision_no=1,
            status="withdrawn",
            submitted_at=now - timedelta(days=2),
            withdrawn_at=now - timedelta(days=1),
        )
        a2_current = _revision(
            request_id="a2-e2-r2",
            student_id="A2",
            course_id="CA",
            version_id=versions[1].id,
            revision_no=2,
            status="submitted",
            submitted_at=now - timedelta(hours=2),
        )
        b1_report = _revision(
            request_id="b1-e1-r1",
            student_id="B1",
            course_id="CB",
            version_id=versions[0].id,
            revision_no=1,
            status="submitted",
            submitted_at=now - timedelta(minutes=30),
            secret="other-teacher-report-secret",
        )
        session.add_all([a1_report, a2_withdrawn, a2_old, a2_current, b1_report])
        session.flush()
        session.add(
            Review(
                submission_id=a1_report.id,
                teacher_id="TA",
                private_score=42,
                private_comment="teacher-private-comment-must-stay-private",
            )
        )
        experiment_ids = [item.id for item in experiments]
    return {"now": now, "experiments": experiment_ids}


def test_teacher_scope_dynamic_tasks_and_zero_submission_students(analytics_app, analytics_data):
    with analytics_app.app_context():
        result = get_teacher_analytics(db_session(), "TA", now=analytics_data["now"])

    assert result["scope"] == {
        "teacher_id": "TA",
        "course_id": "CA",
        "course_name": "甲教师量子实验班",
    }
    assert [course["id"] for course in result["courses"]] == ["CA", "CA-EMPTY"]
    assert result["summary"]["course_experiment_count"] == 3
    assert result["summary"]["task_count"] == 6
    assert result["summary"]["student_count"] == 3
    assert result["summary"]["completed_task_cells"] == 4
    assert result["summary"]["in_progress_task_cells"] == 2
    assert result["summary"]["not_started_task_cells"] == 12
    assert result["summary"]["assignment_count"] == 18
    assert result["summary"]["completion_rate"] == 22.2
    assert result["summary"]["in_progress_rate"] == 11.1
    assert result["summary"]["not_started_rate"] == 66.7
    assert {student["id"] for student in result["students"]} == {"A1", "A2", "A3"}
    zero = next(student for student in result["students"] if student["id"] == "A3")
    assert zero["status"] == "not_started"
    assert all(task["status"] == "not_started" for task in zero["tasks"])
    assert "B1" not in json.dumps(result, ensure_ascii=False)
    assert "乙教师量子实验班" not in json.dumps(result, ensure_ascii=False)


def test_withdrawn_report_is_incomplete_and_resubmission_is_current(analytics_app, analytics_data):
    first_id, second_id = analytics_data["experiments"][:2]
    with analytics_app.app_context():
        result = get_teacher_analytics(db_session(), "TA", now=analytics_data["now"])

    first_report = next(task for task in result["tasks"] if task["id"] == f"{first_id}:report")
    second_report = next(task for task in result["tasks"] if task["id"] == f"{second_id}:report")
    assert first_report["counts"] == {"completed": 1, "in_progress": 1, "not_started": 1}
    assert {student["id"] for student in first_report["completed_students"]} == {"A1"}
    assert {student["id"] for student in first_report["in_progress_students"]} == {"A2"}
    assert second_report["counts"]["completed"] == 1
    current = next(student for student in second_report["completed_students"] if student["id"] == "A2")
    assert current["completed_at"] == (analytics_data["now"] - timedelta(hours=2)).isoformat()


def test_task_status_and_time_filters_are_stable(analytics_app, analytics_data):
    first_id, second_id = analytics_data["experiments"][:2]
    with analytics_app.app_context():
        completed = get_teacher_analytics(
            db_session(),
            "TA",
            experiment_id=first_id,
            task_type="report",
            status="completed",
            now=analytics_data["now"],
        )
        recent = get_teacher_analytics(
            db_session(),
            "TA",
            experiment_id=second_id,
            task_type="report",
            submitted_from=analytics_data["now"] - timedelta(hours=6),
            submitted_to=analytics_data["now"],
            now=analytics_data["now"],
        )

    assert completed["summary"]["task_count"] == 1
    assert completed["summary"]["visible_student_count"] == 1
    assert [student["id"] for student in completed["students"]] == ["A1"]
    assert recent["summary"]["completions_in_range"] == 1
    assert recent["tasks"][0]["completions_in_range"] == 1
    assert sum(point["report"] for point in recent["trends"]) == 1

    with analytics_app.app_context():
        with pytest.raises(ValueError, match="包含时区"):
            get_teacher_analytics(db_session(), "TA", submitted_from=datetime(2026, 7, 1))
        with pytest.raises(ValueError, match="不属于当前教师"):
            get_teacher_analytics(db_session(), "TA", course_id="CB")


def test_empty_scopes_have_stable_shape(analytics_app, analytics_data):
    with analytics_app.app_context():
        no_course = get_teacher_analytics(db_session(), "TC", now=analytics_data["now"])
        empty_course = get_teacher_analytics(
            db_session(), "TA", course_id="CA-EMPTY", now=analytics_data["now"]
        )

    for result in (no_course, empty_course):
        assert result["tasks"] == []
        assert result["students"] == []
        assert result["summary"]["student_count"] == 0
        assert len(result["trends"]) == 7
        assert result["alerts"]


def test_ai_snapshot_is_explicitly_whitelisted_and_private(analytics_app, analytics_data):
    with analytics_app.app_context():
        snapshot = build_teacher_ai_snapshot(db_session(), "TA", now=analytics_data["now"])

    assert set(snapshot) == {
        "scope",
        "generated_at",
        "filters",
        "summary",
        "tasks",
        "students",
        "trends",
        "alerts",
    }
    serialized = json.dumps(snapshot, ensure_ascii=False)
    for forbidden in (
        "report-body-must-stay-private",
        "attachment-hash-must-stay-private",
        "teacher-private-comment-must-stay-private",
        "other-teacher-report-secret",
        "private_score",
        "private_comment",
        "payload",
        "file_hashes",
        "deterministic_score",
        "evaluation_id",
    ):
        assert forbidden not in serialized
    assert "2026001" in serialized and "甲一" in serialized
    assert any(
        task["completed_at"]
        for student in snapshot["students"]
        for task in student["tasks"]
    )


def test_observatory_exposes_scores_warnings_slowest_and_export(analytics_app, analytics_data):
    with analytics_app.app_context():
        result = get_teacher_analytics(db_session(), "TA", now=analytics_data["now"])
        stream = build_teacher_analytics_workbook(result)

    assert result["summary"]["teacher_score_average"] == 42.0
    assert result["summary"]["warning_count"] >= 1
    assert len(result["slowest_students"]) == 3
    assert result["slowest_students"][0]["id"] == "A3"
    assert len(result["experiment_stats"]) == 3
    workbook = load_workbook(io.BytesIO(stream.getvalue()), data_only=True)
    assert workbook.sheetnames == ["班级概览", "学生学情明细", "实验任务完成率"]
    assert workbook["学生学情明细"].max_row == 4


def test_observatory_handles_more_than_one_hundred_students_under_two_seconds(analytics_app, analytics_data):
    with analytics_app.app_context(), transaction() as session:
        for index in range(4, 151):
            student_id = f"A{index}"
            session.add(User(id=student_id, username=f"2026{index:04d}", name=f"性能学生{index}", role="student", password_hash="x"))
            session.add(Enrollment(course_id="CA", student_id=student_id))
    with analytics_app.app_context():
        started = perf_counter()
        result = get_teacher_analytics(db_session(), "TA", now=analytics_data["now"])
        elapsed = perf_counter() - started

    assert result["summary"]["student_count"] == 150
    assert len(result["students"]) == 150
    assert elapsed < 2.0
