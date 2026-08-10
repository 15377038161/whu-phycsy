from __future__ import annotations
import io, json, uuid, zipfile
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from auth_service import password_hash
from ai_service import _validate_question_bank
from platform_app import create_app
from platform_app.content import EXPERIMENTS, definition
from platform_app.db import db_session, transaction
from platform_app.models import AuditEvent, Course, CourseExperiment, Enrollment, Experiment, ExperimentDraft, ExperimentVersion, FileAsset, QuizAssignment, ReportDraft, Review, SubmissionRevision, User
from platform_app.services.scoring import fingerprint


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.delenv("AI_API_KEY", raising=False)
    app = create_app({"TESTING": True, "APP_ENV": "testing", "DATABASE_URL": f"sqlite:///{tmp_path/'test.db'}", "DATA_DIR": str(tmp_path/"files"), "SECRET_KEY": "test"})
    yield app


@pytest.fixture()
def client(app): return app.test_client()


def login(client, username="S001", password="123456"):
    return client.post("/login", data={"username": username, "password": password}, follow_redirects=True)


def test_fingerprint_is_canonical_and_versioned():
    a = fingerprint("v1", {"b": 2, "a": 1}, ["z", "a"])
    b = fingerprint("v1", {"a": 1, "b": 2}, ["a", "z"])
    assert a == b
    assert a != fingerprint("v2", {"a": 1, "b": 2}, ["a", "z"])


def test_login_and_five_official_experiments(client):
    page = login(client)
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    for title in ("离子阱实验", "量子密钥分发", "量子纠缠", "金刚石量子计算机", "单像素光子成像"):
        assert title in text


def test_teacher_ai_experiment_draft_from_voice_text_and_docx(app, client, monkeypatch):
    generated = {
        "code": "HOM",
        "title": "Hong-Ou-Mandel 干涉",
        "definition": {"summary": "面向本科生的双光子干涉实验，通过扫描相对延时测量符合计数凹陷并分析光子不可区分性。", "principle": "两个不可区分单光子同时进入分束器时会发生双光子干涉，理想情况下符合计数在零延时附近下降。", "apparatus": "光子对源、分束器、延时台、单光子探测器与符合计数器。", "formula": "C(tau)=c-a*exp(-(tau/sigma)**2)", "fit_template": "linear", "reference_value": 0.8, "reference_unit": "可见度", "steps": ["关闭探测器高压后检查两路光纤和分束器连接，并记录初始状态。", "开启光子对源并分别测量两路单计数，调节耦合使计数保持稳定。", "设置符合时间窗并扫描电子延时，找到真实符合峰后锁定窗口。", "缓慢扫描光学延时台，每个位置等待稳定后记录单计数与符合计数。", "扣除偶然符合并拟合凹陷曲线，计算可见度和不确定度。"], "questions": [], "report_sections": []},
    }
    monkeypatch.setattr("platform_app.blueprints.teacher.generate_experiment_draft", lambda instruction, material="": generated)
    login(client, "admin", "admin123")
    response = client.post("/teacher/experiments/ai-draft", data={"experiment_prompt": "创建双光子干涉实验", "input_mode": "voice"}, follow_redirects=True)
    html = response.get_data(as_text=True)
    assert response.status_code == 200 and "检查 AI 实验草稿" in html and "Hong-Ou-Mandel 干涉" in html and "待教师确认" in html
    with app.app_context():
        draft = db_session().scalar(select(ExperimentDraft).where(ExperimentDraft.status == "draft"))
        assert draft and draft.source_kind == "voice" and db_session().scalar(select(Experiment).where(Experiment.code == "HOM")) is None
        draft_id = draft.id
    confirmed = client.post(f"/teacher/experiments/ai-draft/{draft_id}/confirm", data={
        "code": generated["code"], "title": generated["title"], "summary": generated["definition"]["summary"], "principle": generated["definition"]["principle"], "apparatus": generated["definition"]["apparatus"], "steps": "\n".join(generated["definition"]["steps"]),
    }, follow_redirects=True)
    assert "AI 实验草稿已创建" in confirmed.get_data(as_text=True)
    with app.app_context():
        draft = db_session().get(ExperimentDraft, draft_id)
        assert draft.status == "confirmed" and draft.experiment_id
        created = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == draft.experiment_id))
        assert created.status == "draft" and len(created.definition["steps"]) == 5

    from docx import Document
    document = Document(); document.add_heading("量子信息实验教学大纲", level=1); document.add_paragraph("课程面向本科生，覆盖量子密钥分发、纠缠测量、数据拟合和实验安全。")
    content = io.BytesIO(); document.save(content); content.seek(0)
    material_generated = {**generated, "code": "HOM2", "title": "材料生成双光子实验"}
    monkeypatch.setattr("platform_app.blueprints.teacher.generate_experiment_draft", lambda instruction, material="": material_generated)
    uploaded = client.post("/teacher/experiments/ai-draft", data={"experiment_prompt": "根据材料起草实验", "experiment_material": (content, "manual.docx")}, content_type="multipart/form-data", follow_redirects=True)
    assert uploaded.status_code == 200 and "检查 AI 实验草稿" in uploaded.get_data(as_text=True)
    with app.app_context():
        material_draft = db_session().scalar(select(ExperimentDraft).where(ExperimentDraft.status == "draft"))
        assert material_draft.source_kind == "docx" and material_draft.source_asset_hash
        assert db_session().get(FileAsset, material_draft.source_asset_hash).original_name == "manual.docx"


def test_student_home_feature_follows_real_learning_progress(client):
    initial = login(client).get_data(as_text=True)
    assert "当前实验 · NEXT EXPERIMENT" in initial and "开始实验" in initial
    assert '<div class="featured-copy">' in initial and "离子阱实验" in initial

    assert client.get("/student/experiment/ENT").status_code == 200
    continued = client.get("/student/home").get_data(as_text=True)
    featured = continued.split('<div class="featured-copy">', 1)[1].split('</div>', 1)[0]
    assert "当前实验 · CONTINUE LEARNING" in featured
    assert "量子纠缠" in featured and "继续实验" in continued


def test_student_hall_lists_every_published_teacher_module(app, client):
    with app.app_context(), transaction() as session:
        source = session.scalar(select(ExperimentVersion).join(Experiment).where(Experiment.code == "ION", ExperimentVersion.status == "published"))
        extra = Experiment(code="NEW", title="教师新增量子模块", order_no=6)
        session.add(extra); session.flush()
        session.add(ExperimentVersion(experiment_id=extra.id, version_no=1, status="published", definition=dict(source.definition)))
    page = login(client).get_data(as_text=True)
    assert "<h2>实验展厅</h2>" in page and "6 个实验展厅" not in page
    assert "教师新增量子模块" in page and 'data-experiment-gallery' in page
    assert 'data-gallery-prev' in page and 'data-gallery-next' in page
    assert page.count('class="experiment-display-card') == 6
    assert client.get("/student/experiment/NEW").status_code == 200


def test_teacher_reorders_experiments_and_student_catalog_follows(app, client):
    login(client, "admin", "admin123")
    page = client.get("/teacher/experiments").get_data(as_text=True)
    assert 'data-experiment-management-list' in page and 'experiment-order-navigation' in page
    assert 'data-experiment-view=' not in page and "teacherExperimentDirectoryView" not in client.get("/static/platform.js").get_data(as_text=True)
    with app.app_context():
        qkd = db_session().scalar(select(Experiment).where(Experiment.code == "QKD"))
        qkd_id = qkd.id
    moved = client.post(f"/teacher/experiments/{qkd_id}/move", data={"direction": "up"}, follow_redirects=True)
    assert "学生端顺序同步更新" in moved.get_data(as_text=True)
    with app.app_context():
        ordered = db_session().scalars(select(Experiment).order_by(Experiment.order_no)).all()
        assert [item.code for item in ordered[:2]] == ["QKD", "ION"]
        assert [item.order_no for item in ordered] == list(range(1, len(ordered) + 1))
        audit = db_session().scalar(select(AuditEvent).where(AuditEvent.action == "experiment.reorder"))
        assert audit and audit.detail == {"direction": "up", "from": 2, "to": 1}
    student = app.test_client(); login(student)
    student_home = student.get("/student/home").get_data(as_text=True)
    assert student_home.index("量子密钥分发") < student_home.index("离子阱实验")


def test_quiz_is_frozen_and_unique_per_student(app, client):
    login(client); assert client.get("/student/experiment/ION").status_code == 200
    with app.app_context(), transaction() as session:
        session.add(User(id="S002", username="S002", name="学生二", role="student", password_hash=password_hash("123456")))
        session.add(Enrollment(course_id="C001", student_id="S002"))
    other = app.test_client(); login(other, "S002"); assert other.get("/student/experiment/ION").status_code == 200
    with app.app_context():
        quizzes = db_session().scalars(select(QuizAssignment).order_by(QuizAssignment.student_id)).all()
        assert len(quizzes) == 2 and quizzes[0].signature != quizzes[1].signature
        original = quizzes[0].questions
        assert len({question.get("concept_id") for question in original}) == 10
    assert client.get("/student/experiment/ION").status_code == 200
    with app.app_context(): assert db_session().scalar(select(QuizAssignment).where(QuizAssignment.student_id == "S001")).questions == original


def test_quiz_feedback_is_frozen_and_retake_excludes_all_original_questions(app, client):
    login(client); client.get("/student/experiment/ION")
    with app.app_context():
        quiz = db_session().scalar(select(QuizAssignment).where(QuizAssignment.student_id == "S001"))
        original_ids = {question["id"] for question in quiz.questions}
        correct_answers = {question["id"]: question["answer"] for question in quiz.questions}
        quiz_id = quiz.id
    submitted = client.post("/student/experiment/ION", data={"action": "quiz", **correct_answers}, follow_redirects=True)
    html = submitted.get_data(as_text=True)
    assert "答对 10/10" in html and html.count("回答正确") == 10 and html.count("解析：") == 10
    overwritten = client.post("/student/experiment/ION", data={"action": "quiz", **{key: "伪造覆盖" for key in correct_answers}}, follow_redirects=True)
    assert "答案不能覆盖" in overwritten.get_data(as_text=True)
    with app.app_context(): assert db_session().get(QuizAssignment, quiz_id).answers == correct_answers
    restarted = client.post("/student/experiment/ION", data={"action": "quiz_retake"}, follow_redirects=True)
    assert "历史原题已全部排除" in restarted.get_data(as_text=True)
    with app.app_context():
        quiz = db_session().get(QuizAssignment, quiz_id)
        assert quiz.attempt_no == 2 and len(quiz.history) == 1 and not quiz.answers and quiz.completed_at is None
        assert original_ids.isdisjoint({question["id"] for question in quiz.questions})


def _answer_quiz(app, client, code="ION"):
    client.get(f"/student/experiment/{code}")
    with app.app_context():
        version = db_session().scalar(select(ExperimentVersion).join(Experiment).where(Experiment.code == code, ExperimentVersion.status == "published"))
        quiz = db_session().scalar(select(QuizAssignment).where(QuizAssignment.student_id == "S001", QuizAssignment.experiment_version_id == version.id))
        data = {"action": "quiz", **{q["id"]: q["answer"] for q in quiz.questions}}
    client.post(f"/student/experiment/{code}", data=data)


def _submit(client, code="ION", request_id=None, value="1.04"):
    return client.post(f"/student/experiment/{code}/submit", data={"request_id": request_id or str(uuid.uuid4()), "abstract": "完整摘要", "principle": "完整实验原理", "raw_data": "1,2\n2,4", "discussion": "误差来自读数和标定。", "conclusion": "完成实验并得到关键结果。", "result_value": value, "steps_complete": "yes", "fit_result": json.dumps({"code": "", "formula": "y=a*x+b", "initial_parameters": {"a": 1}, "final_parameters": {"a": 2}, "metrics": {"r2": .99}})}, follow_redirects=True)


def test_submit_idempotency_has_no_machine_grade_and_teacher_review_stays_private(app, client):
    login(client); _answer_quiz(app, client)
    request_id = str(uuid.uuid4()); first = _submit(client, request_id=request_id); second = _submit(client, request_id=request_id)
    assert "学生端仅展示预习题成绩" in first.get_data(as_text=True)
    with app.app_context():
        rows = db_session().scalars(select(SubmissionRevision)).all()
        assert len(rows) == 1 and rows[0].deterministic_score == 0 and rows[0].relative_error == 0 and rows[0].evaluation_id is None
        revision_id = rows[0].id
    api = client.get("/api/student/submissions").get_json()[0]
    assert set(api).isdisjoint({"private_score", "private_comment", "teacher_score", "deterministic_score", "relative_error", "ai_score", "ai_feedback"})
    teacher = app.test_client(); login(teacher, "admin", "admin123")
    teacher.post(f"/teacher/submission/{revision_id}", data={"private_score": 91, "private_comment": "人工审核"})
    with app.app_context():
        review = db_session().scalar(select(Review).where(Review.submission_id == revision_id))
        assert review and review.private_score == 91


def test_identical_content_keeps_fingerprint_without_creating_ai_evaluation(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context(): first = db_session().scalar(select(SubmissionRevision)); version = db_session().get(ExperimentVersion, first.experiment_version_id)
    client.post(f"/student/submission/{first.id}/withdraw"); _submit(client)
    with app.app_context():
        rows = db_session().scalars(select(SubmissionRevision).order_by(SubmissionRevision.revision_no)).all()
        assert rows[0].fingerprint == rows[1].fingerprint and rows[0].evaluation_id is None and rows[1].evaluation_id is None


def test_two_withdrawals_preserve_reviews_and_third_is_refused(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context(): first = db_session().scalar(select(SubmissionRevision))
    teacher = app.test_client(); login(teacher, "admin", "admin123"); teacher.post(f"/teacher/submission/{first.id}", data={"private_score": 88, "private_comment": "内部记录"})
    client.post(f"/student/submission/{first.id}/withdraw"); _submit(client)
    with app.app_context(): second = db_session().scalar(select(SubmissionRevision).where(SubmissionRevision.revision_no == 2))
    client.post(f"/student/submission/{second.id}/withdraw"); _submit(client)
    with app.app_context():
        third = db_session().scalar(select(SubmissionRevision).where(SubmissionRevision.revision_no == 3)); review = db_session().scalar(select(Review).where(Review.submission_id == first.id)); assert review.superseded
    response = client.post(f"/student/submission/{third.id}/withdraw", follow_redirects=True)
    assert "最多撤回两次" in response.get_data(as_text=True)


def test_teacher_score_never_leaks_to_student_report(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context(): revision = db_session().scalar(select(SubmissionRevision))
    teacher = app.test_client(); login(teacher, "admin", "admin123"); teacher.post(f"/teacher/submission/{revision.id}", data={"private_score": 66, "private_comment": "仅教师可见秘密评语"})
    page = client.get(f"/student/report/{revision.id}").get_data(as_text=True)
    api = client.get("/api/student/submissions").get_data(as_text=True)
    student_surfaces = page + api + client.get("/student/home").get_data(as_text=True) + client.get("/student/records").get_data(as_text=True) + client.get("/student/achievements").get_data(as_text=True)
    assert "仅教师可见秘密评语" not in student_surfaces and "人工评分" not in student_surfaces
    assert "已通过" not in student_surfaces and "奖状" not in student_surfaces and ">66<" not in student_surfaces
    assert "预习题成绩" in student_surfaces and "100/100" in student_surfaces
    internal = teacher.get(f"/teacher/submission/{revision.id}").get_data(as_text=True)
    assert "仅教师可见秘密评语" in internal


def test_student_and_teacher_report_exports_are_separated(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context(): revision = db_session().scalar(select(SubmissionRevision))
    docx = client.get(f"/student/report/{revision.id}.docx"); pdf = client.get(f"/student/report/{revision.id}.pdf")
    assert docx.status_code == 200 and docx.data.startswith(b"PK")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")


def test_report_draft_autosave_rejects_stale_and_unsafe_content(app, client):
    login(client)
    page = client.get("/student/experiment/ION")
    assert "data-report-editor" in page.get_data(as_text=True)
    current = client.get("/student/experiment/ION/report-draft").get_json()
    document = current["content"]
    document["abstract_zh"] = "自动保存后的中文摘要"
    document["fit_result"] = {
        "code": "print('browser-only fit')",
        "formula": "y=a*x+b",
        "initial_parameters": {"a": 1.0, "b": 0.0},
        "final_parameters": {"a": 2.0, "b": 0.1},
        "metrics": {"r2": 0.99},
        "plot_data_url": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZBMkAAAAASUVORK5CYII=",
    }
    raw_section = next(section for section in document["sections"] if section["id"] == "raw_data")
    raw_section["blocks"] = [{"type": "table", "rows": [["x", "y"], ["1", "2"]], "caption": "原始测量数据（拟合输入）"}]
    saved = client.put("/student/experiment/ION/report-draft", json={"lock_version": current["lock_version"], "content": document})
    assert saved.status_code == 200 and saved.get_json()["content"]["abstract_zh"] == "自动保存后的中文摘要"
    assert saved.get_json()["content"]["fit_result"]["plot_data_url"].startswith("data:image/png;base64,")
    assert next(section for section in saved.get_json()["content"]["sections"] if section["id"] == "raw_data")["blocks"][0]["rows"][1] == ["1", "2"]
    stale = client.put("/student/experiment/ION/report-draft", json={"lock_version": current["lock_version"], "content": document})
    assert stale.status_code == 409 and stale.get_json()["code"] == "conflict"
    unsafe = saved.get_json()["content"]
    unsafe["sections"][0]["blocks"] = [{"type": "html", "text": "<script>alert(1)</script>"}]
    rejected = client.put("/student/experiment/ION/report-draft", json={"lock_version": saved.get_json()["lock_version"], "content": unsafe})
    assert rejected.status_code == 400 and "不支持" in rejected.get_json()["error"]


def test_report_draft_submit_freezes_and_copy_creates_next_editable_revision(app, client):
    login(client); _answer_quiz(app, client)
    current = client.get("/student/experiment/ION/report-draft").get_json()
    document = current["content"]
    document.update({"abstract_zh": "结构化报告摘要", "abstract_en": "Structured report abstract.", "keywords_zh": "离子阱；拟合", "keywords_en": "ion trap; fitting", "result_value": "1.04", "steps_complete": True})
    for section in document["sections"]:
        if section["id"] in {"raw_data", "discussion", "conclusion"}:
            section["blocks"] = [{"type": "paragraph", "text": f"{section['title']}内容"}]
    saved = client.put("/student/experiment/ION/report-draft", json={"lock_version": current["lock_version"], "content": document}).get_json()
    submitted = client.post("/student/experiment/ION/submit", data={"request_id": str(uuid.uuid4()), "draft_id": current["id"]}, follow_redirects=True)
    assert submitted.status_code == 200
    with app.app_context():
        revision = db_session().scalar(select(SubmissionRevision)); draft = db_session().get(ReportDraft, current["id"])
        assert revision.payload["report_document"]["abstract_zh"] == "结构化报告摘要" and draft.status == "locked"
        revision_id = revision.id
    copied = client.post(f"/student/submission/{revision_id}/copy-to-draft", headers={"Accept": "application/json"})
    assert copied.status_code == 200 and copied.get_json()["status"] == "active" and copied.get_json()["source_revision_id"] == revision_id
    with app.app_context():
        original = db_session().get(SubmissionRevision, revision_id)
        assert original.revision_no == 1 and original.payload["report_document"]["abstract_zh"] == "结构化报告摘要"


def test_report_draft_exact_pdf_preview(app, client):
    login(client); client.get("/student/experiment/ION")
    current = client.get("/student/experiment/ION/report-draft").get_json()
    preview = client.post("/student/experiment/ION/report-draft/preview.pdf", json={"content": current["content"]})
    assert preview.status_code == 200 and preview.data.startswith(b"%PDF")


def test_uploaded_original_file_is_content_addressed(app, client):
    login(client); _answer_quiz(app, client)
    payload={"request_id":str(uuid.uuid4()),"abstract":"摘要","principle":"原理","raw_data":"1,2","discussion":"讨论","conclusion":"结论","result_value":"1.0","steps_complete":"yes","fit_result":"{}","attachments":(io.BytesIO(b"immutable raw data"),"raw.csv")}
    client.post("/student/experiment/ION/submit",data=payload,content_type="multipart/form-data")
    with app.app_context():
        asset=db_session().scalar(select(FileAsset)); revision=db_session().scalar(select(SubmissionRevision))
        assert asset and asset.sha256 in revision.payload["file_hashes"]
        assert (Path(app.config["DATA_DIR"])/asset.object_key).read_bytes()==b"immutable raw data"


def test_started_experiment_stays_on_pinned_version(app, client):
    login(client); _answer_quiz(app, client)
    with app.app_context(), transaction() as session:
        exp=session.scalar(select(Experiment).where(Experiment.code=="ION")); current=session.scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id==exp.id)); session.add(ExperimentVersion(experiment_id=exp.id,version_no=2,status="published",definition=current.definition))
    page=client.get("/student/experiment/ION").get_data(as_text=True)
    assert "VERSION 1" in page and "VERSION 2" not in page


def test_builtin_question_banks_have_ten_concepts_and_no_irrelevant_options():
    forbidden = ("学生人数", "页面尺寸", "文件名称", "浏览器", "登录时间", "课程完成率", "字体大小")
    for experiment in EXPERIMENTS:
        questions = definition(experiment)["questions"]
        assert len(questions) == 40
        assert len({question["concept_id"] for question in questions}) == 10
        assert not any(word in option for question in questions for option in question["options"] for word in forbidden)


def test_ai_question_quality_gate_rejects_irrelevant_distractors():
    bad = [{"question": f"问题{i}", "options": ["正确实验概念", "页面尺寸", "学生人数", "文件名称"], "answer": "正确实验概念", "explanation": "这是足够长度的实验解析。"} for i in range(40)]
    with pytest.raises(ValueError, match="无关内容"):
        _validate_question_bank(bad, 40)


def test_teacher_roster_has_real_progress_and_working_control_contracts(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    teacher = app.test_client(); login(teacher, "admin", "admin123")
    roster = teacher.get("/teacher/students").get_data(as_text=True)
    dashboard = teacher.get("/teacher").get_data(as_text=True)
    script = (Path(__file__).parents[1] / "static" / "platform.js").read_text(encoding="utf-8")
    assert "1 / 5" in roster and "data-table-course" in roster and "data-course=" in roster
    assert "data-table-sort" in dashboard and "data-sort-value" in dashboard
    assert 'course?.addEventListener("change",filter)' in script
    assert 'sort?.addEventListener("click"' in script


def test_teacher_observatory_has_three_sections_drilldown_and_export(app, client):
    login(client, "admin", "admin123")
    page = client.get("/teacher/analytics?course=C001").get_data(as_text=True)
    assert "班级整体进度概览" in page
    assert "实验任务完成率统计" in page
    assert "学生个体学情明细" in page
    assert "进度最慢的 3 名学生" in page
    assert "data-student-drill" in page and "data-student-details" in page
    assert 'data-analytics-filter-toggle' in page
    assert 'data-student-status-filter="incomplete"' in page
    assert 'data-student-status-filter="warning"' in page
    assert 'data-student-incomplete=' in page and 'data-label="整体进度"' in page
    css = (Path(__file__).parents[1] / "static" / "quantum-platform.css").read_text(encoding="utf-8")
    assert ".analytics-overview-grid .teacher-v4-metrics{grid-template-columns:repeat(4,minmax(0,1fr))}" in css
    exported = client.get("/teacher/analytics.xlsx?course=C001")
    assert exported.status_code == 200 and exported.data.startswith(b"PK")


def test_teacher_information_architecture_has_one_ai_entry_and_class_creation(client):
    login(client, "admin", "admin123")
    dashboard = client.get("/teacher").get_data(as_text=True)
    experiments = client.get("/teacher/experiments").get_data(as_text=True)
    students = client.get("/teacher/students").get_data(as_text=True)
    assert dashboard.count("data-open-teacher-ai") == 1
    assert "teacher-command-summary" not in dashboard and "teacher-course-signal" not in dashboard
    assert "AI 生成实验" in experiments and "手动新建实验" in experiments
    assert "创建班级" in students and "创建课程" not in students
    assert "course-action-panel" not in students and "teacher-v4-metrics" not in students


def test_teacher_experiment_management_is_text_only_and_has_complete_actions(client):
    login(client, "admin", "admin123")
    page = client.get("/teacher/experiments").get_data(as_text=True)
    experiment_list = page.split('data-experiment-management-list', 1)[1].split('</section>', 1)[0]
    assert "实验管理" in page and "experiment-management-row" in experiment_list
    assert "experiment-admin-cover" not in experiment_list and "admin-version-stack" not in experiment_list
    assert "版本时间线" not in experiment_list and "<img" not in experiment_list
    for action in ("编辑实验", "复制实验", "发布实验", "下架实验", "删除实验"):
        assert action in experiment_list


def test_teacher_copies_experiment_as_independent_draft(app, client):
    login(client, "admin", "admin123")
    with app.app_context():
        source = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
        source_id = source.id
    response = client.post(f"/teacher/experiments/{source_id}/copy", follow_redirects=False)
    assert response.status_code == 302 and "/draft" in response.headers["Location"]
    with app.app_context():
        copied = db_session().scalar(select(Experiment).where(Experiment.code.like("ION_COPY%")))
        draft = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == copied.id))
        links = db_session().scalars(select(CourseExperiment).where(CourseExperiment.experiment_id == copied.id)).all()
        audit = db_session().scalar(select(AuditEvent).where(AuditEvent.action == "experiment.copy", AuditEvent.entity_id == copied.id))
        assert copied.id != source_id and copied.title.startswith("离子阱实验（副本")
        assert draft.status == "draft" and draft.version_no == 1
        assert links and audit.detail["source_experiment_id"] == source_id


def test_experiment_editor_has_back_previous_and_next_navigation(client):
    login(client, "admin", "admin123")
    with client.application.app_context():
        experiment = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
    editor = client.get(f"/teacher/experiments/{experiment.id}/draft").get_data(as_text=True)
    script = (Path(__file__).parents[1] / "static" / "platform.js").read_text(encoding="utf-8")
    assert "返回实验管理" in editor and "data-editor-flow-nav" in editor
    assert "data-editor-previous" in editor and "data-editor-next" in editor
    assert 'previous?.addEventListener("click"' in script and 'next?.addEventListener("click"' in script


def test_experiment_editor_uses_teacher_friendly_question_manager_and_light_surfaces(client):
    login(client, "admin", "admin123")
    with client.application.app_context():
        experiment = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
    editor = client.get(f"/teacher/experiments/{experiment.id}/draft").get_data(as_text=True)
    script = (Path(__file__).parents[1] / "static" / "platform.js").read_text(encoding="utf-8")
    assert "候选题 JSON（高级编辑）" not in editor and 'class="code-editor"' not in editor
    assert "题目审核与管理" in editor and "data-question-manager" in editor
    assert "data-question-list" in editor and "data-question-form" in editor
    assert "data-question-add" in editor and 'name="questions_json" hidden' in editor
    assert 'name="mode" value="append"' in editor and 'name="mode" value="replace"' in editor
    assert "AI 新增题目" in editor and "全部重新生成" in editor
    assert "experiment-editor-hero" in editor and "experiment-editor-aside" in editor
    assert "function setupQuestionManager()" in script and "reviewed_by_teacher=true" in script


def test_platform_ui_uses_one_typography_icon_and_page_scale_standard():
    root = Path(__file__).parents[1]
    css = (root / "static" / "quantum-platform.css").read_text(encoding="utf-8")
    guide = (root / "UI_VISUAL_STANDARD.md").read_text(encoding="utf-8")
    templates = "\n".join(path.read_text(encoding="utf-8") for path in (root / "templates").glob("*.html"))
    for token in ("--ui-font-body", "--ui-font-display", "--ui-title-page", "--ui-text-body", "--ui-icon-control", "--ui-content-wide"):
        assert token in css
    assert "1328 px" in guide and "40 px" in guide and "16 × 16 px" in guide
    assert "font-size:" not in templates and "font-family:" not in templates
    assert "20260806-ui-standard-v1" in (root / "templates" / "base.html").read_text(encoding="utf-8")


def test_first_password_page_uses_the_global_visual_system(app, client):
    with app.app_context(), transaction() as tx:
        student = tx.scalar(select(User).where(User.username == "S001"))
        student.must_change_password = True
    response = client.post("/login", data={"username": "S001", "password": "123456"}, follow_redirects=True)
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "首次登录，请设置新密码" in page
    assert 'class="first-password-page page-pad"' in page and 'class="first-password-card"' in page


def test_teacher_can_append_or_replace_ai_questions_without_accidental_overwrite(app, client, monkeypatch):
    login(client, "admin", "admin123")
    with app.app_context():
        experiment = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
        client.get(f"/teacher/experiments/{experiment.id}/draft")
        draft = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == experiment.id, ExperimentVersion.status == "draft"))
        original_count = len(draft.definition.get("questions", []))
        source_questions = deepcopy(draft.definition.get("questions", []))
    generated = []
    for index in range(40):
        base = deepcopy(source_questions[index % len(source_questions)])
        base.update({"id": f"generated-{index + 1}", "question": f"AI 新题 {index + 1}：{base['question']}"})
        generated.append(base)
    monkeypatch.setattr("platform_app.blueprints.teacher.generate_question_bank", lambda title, corpus, count: deepcopy(generated))
    appended = client.post(f"/teacher/experiments/version/{draft.id}/generate-questions", data={"mode": "append", "count": "3"}, follow_redirects=True)
    assert appended.status_code == 200 and "已新增 3 道候选题" in appended.get_data(as_text=True)
    with app.app_context():
        updated = db_session().get(ExperimentVersion, draft.id)
        assert len(updated.definition["questions"]) == original_count + 3
        assert updated.definition["questions"][-1]["reviewed_by_teacher"] is False
    replaced = client.post(f"/teacher/experiments/version/{draft.id}/generate-questions", data={"mode": "replace", "count": "40"}, follow_redirects=True)
    assert replaced.status_code == 200 and "已重新生成 40 道候选题" in replaced.get_data(as_text=True)
    with app.app_context():
        updated = db_session().get(ExperimentVersion, draft.id)
        audit = db_session().scalar(select(AuditEvent).where(AuditEvent.action == "questions.generate", AuditEvent.entity_id == draft.id).order_by(AuditEvent.created_at.desc()))
        assert len(updated.definition["questions"]) == 40
        assert audit.detail["mode"] == "replace" and audit.detail["total_count"] == 40


def test_teacher_mobile_review_places_scoring_before_report_and_has_quick_navigation(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context():
        revision = db_session().scalar(select(SubmissionRevision))
    teacher = app.test_client(); login(teacher, "admin", "admin123")
    page = teacher.get(f"/teacher/submission/{revision.id}").get_data(as_text=True)
    assert 'class="teacher-mobile-review-nav page-pad"' in page
    assert 'id="teacherReviewPanel"' in page and 'id="reportContent"' in page
    assert page.index('id="reportContent"') < page.index('id="teacherReviewPanel"')
    assert 'inputmode="numeric"' in page and 'class="teacher-number" name="private_score"' in page


def test_teacher_can_delete_draft_only_experiment_and_safely_archive_published_one(app, client):
    login(client, "admin", "admin123")
    created = client.post("/teacher/experiments/new", data={"code": "TMP", "title": "临时实验"}, follow_redirects=True)
    assert created.status_code == 200
    with app.app_context():
        draft_only = db_session().scalar(select(Experiment).where(Experiment.code == "TMP"))
        published = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
        draft_only_id, published_id = draft_only.id, published.id
    deleted = client.post(f"/teacher/experiments/{draft_only_id}/delete", follow_redirects=True)
    assert "未发布实验及其草稿版本已彻底删除" in deleted.get_data(as_text=True)
    archived = client.post(f"/teacher/experiments/{published_id}/delete", follow_redirects=True)
    assert "已安全删除出学生端并归档" in archived.get_data(as_text=True)
    with app.app_context():
        assert db_session().get(Experiment, draft_only_id) is None
        published = db_session().get(Experiment, published_id)
        versions = db_session().scalars(select(ExperimentVersion).where(ExperimentVersion.experiment_id == published_id)).all()
        assert published.retired is True
        assert versions and all(version.status != "published" for version in versions)


def test_student_operation_tab_is_detailed_text_not_fake_progress(client):
    login(client)
    page = client.get("/student/experiment/ION").get_data(as_text=True)
    assert "实验操作步骤" in page and "operation-procedure" in page
    assert "当前步骤" not in page and "待开始" not in page


def test_fitting_table_has_no_row_action_column_and_uses_add_row_label(client):
    login(client)
    page = client.get("/student/experiment/QKD").get_data(as_text=True)
    table = page.split('<table class="fit-table">', 1)[1].split("</table>", 1)[0]
    assert "添加行" in page and "添加测量点" not in page
    assert "row-remove" not in table and "删除此行" not in table
    header = table.split("</thead>", 1)[0]
    assert header.count("<th>") == 3


def test_student_home_uses_100_point_quiz_score_and_completed_experiment_count(app, client):
    login(client)
    _answer_quiz(app, client)
    before_submission = client.get("/student/home").get_data(as_text=True)
    assert "预习题成绩 <b>100/100</b>" in before_submission
    assert "已完成实验" in before_submission and "0 / 5" in before_submission
    _submit(client)
    after_submission = client.get("/student/home").get_data(as_text=True)
    assert "已完成实验" in after_submission and "1 / 5" in after_submission


def test_learning_memorial_unlocks_from_submission_only_and_relocks_after_withdrawal(app, client):
    login(client)
    initial = client.get("/student/achievements").get_data(as_text=True)
    assert "荣誉展厅" in initial and initial.count("data-achievement-item") == 5
    assert initial.count('data-unlocked="false"') == 5
    assert "正式提交实验报告后解锁荣誉证书" in initial and "通过后解锁" not in initial

    _answer_quiz(app, client)
    _submit(client)
    unlocked = client.get("/student/achievements").get_data(as_text=True)
    ion = unlocked.split('data-code="ION"', 1)[1].split("</button>", 1)[0]
    assert 'data-unlocked="true"' in ion and 'data-revision="1"' in ion
    assert "实验探索纪念证书" in unlocked and "完成全部实验流程并正式提交报告" in unlocked

    with app.app_context():
        revision = db_session().scalar(select(SubmissionRevision).where(SubmissionRevision.status == "submitted"))
        revision_id = revision.id
    client.post(f"/student/submission/{revision_id}/withdraw")
    relocked = client.get("/student/achievements").get_data(as_text=True)
    ion = relocked.split('data-code="ION"', 1)[1].split("</button>", 1)[0]
    assert 'data-unlocked="false"' in ion


def test_student_experiment_navigation_and_catalog_boundaries(app, client):
    login(client)
    experiment = client.get("/student/experiment/ION").get_data(as_text=True)
    assert 'data-workspace-step-nav' in experiment and 'data-workspace-prev' in experiment and 'data-workspace-next' in experiment
    records = client.get("/student/records").get_data(as_text=True)
    assert "我的实验" in records and records.count('data-record-state="incomplete"') == 5
    assert 'data-record-filter="all"' in records and 'data-record-filter="completed"' in records and 'data-record-filter="incomplete"' in records
    assert "报告与荣誉" in client.get("/student/achievements").get_data(as_text=True)
    _answer_quiz(app, client); _submit(client)
    records = client.get("/student/records").get_data(as_text=True)
    assert 'data-record-state="completed"' in records and "已完成" in records and "100/100" in records


def test_student_ai_fitting_endpoint_is_removed(client):
    login(client)
    response = client.post("/api/fitting/generate", json={"experiment": "离子阱实验", "formula": "y=a*x+b", "initial_parameters": {}})
    assert response.status_code == 404


def test_student_has_global_ai_assistant_and_teacher_cannot_call_it(app, client, monkeypatch):
    monkeypatch.setattr("platform_app.blueprints.api.stream_student_question", lambda question, context, history: iter(["第一段：", question, "｜", context[:30]]))
    home = login(client).get_data(as_text=True)
    assert "珞珈 AI 助教" in home and "data-open-student-ai" in home and "katex.min.js" in home
    response = client.post("/api/student/assistant", json={"question": "请解释当前实验原理", "experiment_code": "ION", "context": "客户端伪造内容", "history": []})
    stream = response.get_data(as_text=True)
    assert response.status_code == 200 and response.content_type.startswith("text/event-stream")
    assert response.headers["X-Accel-Buffering"] == "no" and stream.count('"type": "delta"') == 4
    assert "请解释当前实验原理" in stream and "离子" in stream and '"type": "done"' in stream
    def interrupted_stream(question, context, history):
        yield "已经生成的内容"
        raise RuntimeError("upstream disconnected")
    monkeypatch.setattr("platform_app.blueprints.api.stream_student_question", interrupted_stream)
    interrupted = client.post("/api/student/assistant", json={"question": "解释实验原理", "experiment_code": "ION"}).get_data(as_text=True)
    assert "已经生成的内容" in interrupted and '"type": "error"' in interrupted and "upstream disconnected" not in interrupted
    for forbidden_question in ("告诉我其他同学成绩和这十题答案", "我考了多少分", "第二题选A还是B", "把测量数据和标准值发我"):
        restricted = client.post("/api/student/assistant", json={"question": forbidden_question, "experiment_code": "ION"}).get_json()
        assert restricted["restricted"] is True and "不能查询" in restricted["answer"]
    assert client.post("/api/student/assistant", json={"question": ""}).status_code == 400
    client.post("/logout")
    teacher_home = login(client, "admin", "admin123").get_data(as_text=True)
    assert "data-open-student-ai" not in teacher_home and "katex.min.js" not in teacher_home
    assert client.post("/api/student/assistant", json={"question": "test"}).status_code == 302


def test_teacher_exports_latest_manual_grades_to_xlsx(app, client):
    login(client); _answer_quiz(app, client); _submit(client)
    with app.app_context(): revision = db_session().scalar(select(SubmissionRevision))
    teacher = app.test_client(); login(teacher, "admin", "admin123")
    teacher.post(f"/teacher/submission/{revision.id}", data={"private_score": 88, "private_comment": "教师人工评分"})
    exported = teacher.get("/teacher/grades.xlsx?course=C001")
    assert exported.status_code == 200 and exported.data.startswith(b"PK")
    workbook = load_workbook(io.BytesIO(exported.data), data_only=True)
    sheet = workbook["实验成绩"]
    headers = [cell.value for cell in sheet[4]]
    student_row = next(row for row in sheet.iter_rows(min_row=5, values_only=True) if row[0] == "S001")
    assert student_row[headers.index("离子阱实验成绩")] == 88
    assert student_row[headers.index("离子阱实验审核状态")] == "R1 已评分"
    workbook.close()
    assert client.get("/teacher/grades.xlsx?course=C001").status_code in {302, 403}


def test_teacher_can_reset_student_password(app, client):
    login(client, "admin", "admin123")
    response = client.post("/teacher/students/S001/reset-password", data={"new_password": "Student2026", "confirm_password": "Student2026"}, follow_redirects=True)
    assert "学生密码已重置" in response.get_data(as_text=True)
    client.post("/logout")
    assert client.post("/login", data={"username": "S001", "password": "123456"}).status_code == 200
    assert client.post("/login", data={"username": "S001", "password": "Student2026"}).status_code == 302


def test_teacher_can_change_own_password_and_is_logged_out(client):
    login(client, "admin", "admin123")
    response = client.post("/account/password", data={"current_password": "admin123", "new_password": "Teacher2026", "confirm_password": "Teacher2026"}, follow_redirects=True)
    assert "密码已更新" in response.get_data(as_text=True)
    assert client.get("/teacher").status_code == 302
    assert client.post("/login", data={"username": "admin", "password": "Teacher2026"}).status_code == 302


def test_teacher_can_save_publish_delete_and_optionally_upload_video(app, client):
    login(client, "admin", "admin123")
    with app.app_context():
        exp = db_session().scalar(select(Experiment).where(Experiment.code == "ION"))
    assert client.get(f"/teacher/experiments/{exp.id}/draft").status_code == 200
    with app.app_context():
        draft = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == exp.id, ExperimentVersion.status == "draft"))
        definition = draft.definition
    payload = {
        "title": "离子阱实验", "summary": definition["summary"], "principle": definition["principle"],
        "apparatus": definition["apparatus"], "reference_value": str(definition["reference_value"]),
        "reference_unit": definition["reference_unit"], "formula": definition["formula"],
        "steps": "\n".join(definition["steps"]), "questions_json": json.dumps(definition["questions"], ensure_ascii=False),
        "intent": "publish", "demo_video": (io.BytesIO(b"\x00\x00\x00\x18ftypmp42demo"), "operation-demo.mp4"),
        "cover_image": (io.BytesIO(b"\x89PNG\r\n\x1a\ncover"), "ion-cover.png"),
    }
    published = client.post(f"/teacher/experiments/{exp.id}/draft", data=payload, content_type="multipart/form-data", follow_redirects=True)
    assert "草稿已保存、发布并锁定" in published.get_data(as_text=True)
    with app.app_context():
        version = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == exp.id).order_by(ExperimentVersion.version_no.desc()))
        assert version.status == "published" and version.definition["demo_video_hash"] and version.definition["cover_image_hash"]
        published_version_id, published_version_no = version.id, version.version_no
        video_hash = version.definition["demo_video_hash"]
        cover_hash = version.definition["cover_image_hash"]
    media = client.get(f"/media/{video_hash}")
    assert media.status_code == 200 and media.data.endswith(b"demo")
    cover = client.get(f"/media/{cover_hash}")
    assert cover.status_code == 200 and cover.data.startswith(b"\x89PNG")
    client.get(f"/teacher/experiments/{exp.id}/draft")
    with app.app_context():
        new_draft = db_session().scalar(select(ExperimentVersion).where(ExperimentVersion.experiment_id == exp.id, ExperimentVersion.status == "draft"))
    deleted = client.post(f"/teacher/experiments/version/{new_draft.id}/delete", follow_redirects=True)
    assert "草稿已删除" in deleted.get_data(as_text=True)
    retired = client.post(f"/teacher/experiments/version/{published_version_id}/retire", follow_redirects=True)
    assert "版本已从学生端下架" in retired.get_data(as_text=True)
    with app.app_context():
        assert db_session().get(ExperimentVersion, published_version_id).status == "retired"
    student = app.test_client(); hall = login(student).get_data(as_text=True)
    assert "离子阱实验" in hall and f"版本 {published_version_no}" not in hall


def test_production_requires_postgresql(tmp_path):
    with pytest.raises(RuntimeError, match="PostgreSQL"):
        create_app({"APP_ENV": "production", "DATABASE_URL": f"sqlite:///{tmp_path/'bad.db'}"})
