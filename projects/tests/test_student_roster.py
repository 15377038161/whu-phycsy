from __future__ import annotations

import io

import pytest
from openpyxl import Workbook
from sqlalchemy import select
from werkzeug.datastructures import FileStorage

from platform_app import create_app
from platform_app.db import db_session
from platform_app.models import AuditEvent, Course, Enrollment, User
from platform_app.services.student_roster import parse_student_roster


def xlsx_upload(sheets: list[tuple[str, list[list[object]]]], filename: str = "学生名单.xlsx") -> FileStorage:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return FileStorage(
        stream=stream,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.fixture
def roster_app(tmp_path, monkeypatch):
    monkeypatch.delenv("AI_API_KEY", raising=False)
    return create_app({
        "TESTING": True,
        "APP_ENV": "testing",
        "DATABASE_URL": f"sqlite:///{tmp_path / 'roster.db'}",
        "DATA_DIR": str(tmp_path / "files"),
        "SECRET_KEY": "roster-test",
    })


def test_parser_finds_shifted_header_and_arbitrary_columns():
    upload = xlsx_upload([
        ("说明", [["学生名单导入说明"], ["本页没有学生数据"]]),
        ("2026名单", [
            ["2026 量子实验班学生花名册"],
            [],
            ["学生姓名（必填）", "专业方向", "校内学籍号", "手机号", "备注"],
            ["张三", "物理学", "20260001", "13800000001", "班长"],
            ["李四", "光电", "20260002", "13800000002", ""],
        ]),
    ])
    parsed = parse_student_roster(upload, "Student2026")
    assert parsed.sheet_name == "2026名单"
    assert parsed.header_row == 3
    assert parsed.student_id_column == "校内学籍号"
    assert parsed.name_column == "学生姓名（必填）"
    assert parsed.ignored_columns == 3
    assert [(row.student_id, row.name, row.password) for row in parsed.rows] == [
        ("20260001", "张三", "Student2026"),
        ("20260002", "李四", "Student2026"),
    ]


def test_parser_uses_data_features_when_id_header_is_unfamiliar_and_name_is_third_column():
    upload = xlsx_upload([("Sheet1", [
        ["校内唯一标识", "学院备注", "中文姓名"],
        ["STU-2026-001", "物理学院", "王小明"],
        ["STU-2026-002", "物理学院", "赵小红"],
        ["STU-2026-003", "物理学院", "陈志远"],
    ])])
    parsed = parse_student_roster(upload, "Student2026")
    assert parsed.student_id_column == "校内唯一标识"
    assert parsed.name_column == "中文姓名"
    assert [row.name for row in parsed.rows] == ["王小明", "赵小红", "陈志远"]


def test_parser_preserves_zero_padded_numeric_student_ids():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "名单"
    sheet.append(["姓名", "学号"])
    sheet.append(["张三", 1234])
    sheet.append(["李四", 5678])
    for cell in (sheet["B2"], sheet["B3"]):
        cell.number_format = "00000000"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    parsed = parse_student_roster(FileStorage(stream=stream, filename="名单.xlsx"), "Student2026")
    assert [row.student_id for row in parsed.rows] == ["00001234", "00005678"]


def test_parser_rejects_duplicate_or_ambiguous_student_ids():
    duplicate = xlsx_upload([("名单", [
        ["姓名", "学号"], ["张三", "20260001"], ["李四", "20260001"],
    ])])
    with pytest.raises(ValueError, match="重复"):
        parse_student_roster(duplicate, "Student2026")

    ambiguous = xlsx_upload([("名单", [
        ["字段甲", "字段乙", "姓名"],
        ["20260001", "A20260001", "张三"],
        ["20260002", "A20260002", "李四"],
        ["20260003", "A20260003", "王五"],
    ])])
    with pytest.raises(ValueError, match="多个可能"):
        parse_student_roster(ambiguous, "Student2026")


def test_teacher_imports_excel_with_row_passwords_and_extra_fields(roster_app):
    client = roster_app.test_client()
    client.post("/login", data={"username": "admin", "password": "admin123"})
    with roster_app.app_context():
        teacher = db_session().scalar(select(User).where(User.username == "admin"))
        course = db_session().scalar(select(Course).where(Course.teacher_id == teacher.id))
        course_id = course.id
    upload = xlsx_upload([("花名册", [
        ["学院", "临时密码", "姓名", "手机号", "学号"],
        ["物理学院", "RowPass2026", "测试甲", "13800000001", "20269901"],
        ["物理学院", "RowPass2027", "测试乙", "13800000002", "20269902"],
    ])])
    response = client.post(
        "/teacher/students",
        data={"action": "students", "course_id": course_id, "student_roster": (upload.stream, upload.filename)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "已从“花名册”识别 2 名学生" in page
    assert "学号列“学号”、姓名列“姓名”" in page
    assert "忽略 2 个其他字段" in page
    with roster_app.app_context():
        students = db_session().scalars(select(User).where(User.id.in_(["20269901", "20269902"]))).all()
        enrollments = db_session().scalars(select(Enrollment).where(Enrollment.course_id == course_id, Enrollment.student_id.in_(["20269901", "20269902"]))).all()
        audit = db_session().scalar(select(AuditEvent).where(AuditEvent.action == "students.import").order_by(AuditEvent.created_at.desc()))
    assert {student.name for student in students} == {"测试甲", "测试乙"}
    assert len(enrollments) == 2
    assert audit.detail["student_id_column"] == "学号"
    assert audit.detail["password_column"] == "临时密码"
    client.post("/logout")
    assert client.post("/login", data={"username": "20269901", "password": "RowPass2026"}).status_code == 302
